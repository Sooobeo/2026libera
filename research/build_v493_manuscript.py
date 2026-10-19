"""Build the progress manuscript without changing frozen experiment artifacts."""
from pathlib import Path
import json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'paper/v493_manuscript_20261019.json'
OUT = ROOT / 'paper/LIBERA_소논문_연구설계와진행현황_20261019.docx'

def set_font(style, size, bold=False):
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii','hAnsi','eastAsia','cs'):
        rf.set(qn('w:'+key), 'Malgun Gothic')

def hyperlink(p, text, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'000000'); pr.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'),'single'); pr.append(under)
    run.append(pr); txt = OxmlElement('w:t'); txt.text = text; run.append(txt)
    link.append(run); p._p.append(link)

def add_table(doc, block):
    cap = doc.add_paragraph(block['caption'], 'Caption')
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(block['headers']))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, width in zip(t.columns, block['widths']): col.width = Cm(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        edge=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: edge.set(qn('w:'+k),v)
        borders.append(edge)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,value in [('top',80),('bottom',80),('left',105),('right',105)]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for i,text in enumerate(block['headers']): t.rows[0].cells[i].text=text
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in block['rows']:
        cells=t.add_row().cells
        for c,text in zip(cells,row):c.text=str(text)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(block['widths'][ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.style=doc.styles['Table Text']
                p.paragraph_format.keep_with_next=False
                if ri==0 or (block['caption'].startswith(('표 1','표 5')) and ci!=0) or (block['caption'].startswith('표 6') and ci==1):
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.bold=(ri==0)
    note=doc.add_paragraph(block.get('note',''),'Table Note')
    note.paragraph_format.keep_with_next=False

def main():
    m=json.loads(SOURCE.read_text(encoding='utf-8'))
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0)
    sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    sec.footer_distance=Cm(1)
    for name,size,bold in [('Normal',11,False),('Title',19,True),('Subtitle',12,False),('Heading 1',14,True),('Heading 2',11.5,True),('Caption',10,True)]:
        set_font(doc.styles[name],size,bold)
    for name,size in [('Table Text',9.5),('Table Note',9),('Reference',9.5)]:
        st=doc.styles.add_style(name,1);st.base_style=doc.styles['Normal'];set_font(st,size)
    norm=doc.styles['Normal'].paragraph_format
    norm.line_spacing=1.28;norm.space_after=Pt(7);norm.widow_control=True
    for name in ('Heading 1','Heading 2'):
        pf=doc.styles[name].paragraph_format
        pf.keep_with_next=True;pf.keep_together=True;pf.space_before=Pt(15);pf.space_after=Pt(7)
    for name in ('Table Text','Table Note','Reference'):
        pf=doc.styles[name].paragraph_format;pf.line_spacing=1.15;pf.space_after=Pt(0 if name=='Table Text' else 6)
    doc.styles['Reference'].paragraph_format.left_indent=Cm(.6)
    doc.styles['Reference'].paragraph_format.first_line_indent=Cm(-.6)
    doc.styles['Table Note'].paragraph_format.space_before=Pt(4)
    # Clear inherited decoration and theme colours from all title/heading styles.
    for name in ('Title','Subtitle','Heading 1','Heading 2','Caption'):
        el=doc.styles[name].element
        for border in el.xpath('.//w:pBdr'):border.getparent().remove(border)
        for color in el.xpath('.//w:color'):
            color.attrib.clear();color.set(qn('w:val'),'000000')
    title=doc.add_paragraph(m['title'],'Title');title.paragraph_format.space_after=Pt(7)
    sub=doc.add_paragraph(m['subtitle'],'Subtitle');sub.paragraph_format.space_after=Pt(10)
    meta=doc.add_paragraph(m['date']+'  |  '+m['version']);meta.paragraph_format.space_after=Pt(15)
    for r in meta.runs:r.font.size=Pt(9)
    for b in m['blocks']:
        if b['type']=='heading':
            p=doc.add_paragraph(b['text'],'Heading '+str(b['level']))
            if b['text']=='참고문헌':p.paragraph_format.page_break_before=True
        elif b['type']=='paragraph':doc.add_paragraph(b['text'])
        elif b['type']=='table':add_table(doc,b)
        elif b['type']=='reference':
            p=doc.add_paragraph(b['text']+' ','Reference')
            if b.get('url'):hyperlink(p,'원문',b['url'])
            for i,link in enumerate(b.get('links',[])):
                if i:p.add_run('  ')
                hyperlink(p,link['label'],link['url'])
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=foot.add_run();r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    doc.core_properties.title=m['title']
    doc.core_properties.subject=m['subtitle']
    doc.core_properties.author=''
    doc.core_properties.keywords='LIBERA, 생성형 AI, 산업배분, 공개 투자판단'
    doc.core_properties.comments=''
    doc.save(OUT)
    receipt={'source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output':str(OUT),'output_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'blocks':len(m['blocks']),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'frozen_experiment_files_modified':False}
    (ROOT/'research/v493_docx_build_20261019.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(receipt,ensure_ascii=False))

if __name__=='__main__':main()
