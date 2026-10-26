"""Build the updated paper DOCX and HTML from one reviewed JSON source."""
from pathlib import Path
import json, hashlib, html, re, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'paper/v500_manuscript_20261026.json'
DOCX = ROOT / 'paper/LIBERA_소논문_실험결과와해결방향_v5.0_20261026.docx'
HTML = ROOT / 'paper/LIBERA_텀페이퍼_v5.0_20261026.html'
QA = ROOT / 'research/paper_update_20261026/qa'

def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()

def set_font(style, size, bold=False):
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+key), 'Malgun Gothic')

def hyperlink(p, text, url):
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'000000'); pr.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'),'single'); pr.append(under)
    run.append(pr); txt=OxmlElement('w:t');txt.text=text;run.append(txt);link.append(run);p._p.append(link)

def table(doc, b):
    p=doc.add_paragraph(b['caption'],'Caption');p.paragraph_format.keep_with_next=True
    n=len(b['headers']); widths=b.get('widths',[16/n]*n)
    if abs(sum(widths)-16)>.05: widths=[w*16/sum(widths) for w in widths]
    t=doc.add_table(rows=1,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        edge=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:edge.set(qn('w:'+k),v)
        borders.append(edge)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,value in [('top',85),('bottom',85),('left',105),('right',105)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for i,text in enumerate(b['headers']):t.rows[0].cells[i].text=str(text)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for vals in b['rows']:
        if len(vals)!=n:raise ValueError('Row width mismatch '+b['caption'])
        for c,text in zip(t.add_row().cells,vals):c.text=str(text)
    align=b.get('align',[])
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF3');cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.style=doc.styles['Table Text'];p.paragraph_format.keep_with_next=(ri==0)
                if ri==0 or (ci<len(align) and align[ci]=='center') or (ri>0 and re.fullmatch(r'[\d.,%+−\- ]+',cell.text)):
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.bold=ri==0
    p=doc.add_paragraph(b.get('note',''),'Table Note');p.paragraph_format.keep_with_next=False

def main():
    m=json.loads(SOURCE.read_text(encoding='utf-8-sig'))
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5);sec.footer_distance=Cm(1)
    for name,size,bold in [('Normal',11,False),('Title',19,True),('Subtitle',12,False),('Heading 1',14,True),('Heading 2',11.5,True),('Heading 3',11,True),('Caption',10,True)]:set_font(doc.styles[name],size,bold)
    for name,size in [('Table Text',9.5),('Table Note',9),('Reference',9.5)]:
        st=doc.styles.add_style(name,1);st.base_style=doc.styles['Normal'];set_font(st,size)
    pf=doc.styles['Normal'].paragraph_format;pf.line_spacing=1.28;pf.space_after=Pt(7);pf.widow_control=True
    for name in ('Heading 1','Heading 2','Heading 3'):
        pf=doc.styles[name].paragraph_format;pf.keep_with_next=True;pf.keep_together=True;pf.space_before=Pt(15);pf.space_after=Pt(7)
    for name in ('Table Text','Table Note','Reference'):
        pf=doc.styles[name].paragraph_format;pf.line_spacing=1.15;pf.space_after=Pt(0 if name=='Table Text' else 6)
    doc.styles['Reference'].paragraph_format.left_indent=Cm(.6);doc.styles['Reference'].paragraph_format.first_line_indent=Cm(-.6)
    doc.styles['Table Note'].paragraph_format.space_before=Pt(4)
    for name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
        el=doc.styles[name].element
        for border in el.xpath('.//w:pBdr'):border.getparent().remove(border)
        for color in el.xpath('.//w:color'):color.attrib.clear();color.set(qn('w:val'),'000000')
    doc.add_paragraph(m['title'],'Title').paragraph_format.space_after=Pt(7)
    doc.add_paragraph(m.get('subtitle',''),'Subtitle').paragraph_format.space_after=Pt(10)
    meta=doc.add_paragraph(m.get('date','')+'  |  '+m.get('version',''));meta.paragraph_format.space_after=Pt(15)
    for r in meta.runs:r.font.size=Pt(9)
    esc=lambda s:html.escape(str(s))
    content=[f'<header><h1>{esc(m["title"])}</h1><p class="subtitle">{esc(m.get("subtitle",""))}</p><p class="meta">{esc(m.get("date",""))} · {esc(m.get("version",""))}</p></header>']
    headings=[]
    for i,b in enumerate(m['blocks']):
        typ=b['type']
        if typ=='heading':
            lev=b['level'];p=doc.add_paragraph(b['text'],'Heading '+str(lev))
            if b['text']=='참고문헌':p.paragraph_format.page_break_before=True
            content.append(f'<h{lev+1} id="s{i}">{esc(b["text"])}</h{lev+1}>')
            headings.append({'level':lev,'text':b['text'],'id':'s'+str(i)})
        elif typ=='paragraph':
            doc.add_paragraph(b['text']);content.append('<p>'+esc(b['text'])+'</p>')
        elif typ=='table':
            table(doc,b)
            content.append('<figure><figcaption>'+esc(b['caption'])+'</figcaption><div class="table-wrap"><table><thead><tr>'+''.join('<th>'+esc(x)+'</th>' for x in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(x)+'</td>' for x in row)+'</tr>' for row in b['rows'])+'</tbody></table></div><p class="note">'+esc(b.get('note',''))+'</p></figure>')
        elif typ=='reference':
            p=doc.add_paragraph(b['text']+' ','Reference');hrefs=[]
            if b.get('url'):hyperlink(p,'원문',b['url']);hrefs.append('<a href="'+esc(b['url'])+'">원문</a>')
            for link in b.get('links',[]):
                p.add_run('  ');hyperlink(p,link['label'],link['url']);hrefs.append('<a href="'+esc(link['url'])+'">'+esc(link['label'])+'</a>')
            content.append('<p class="reference">'+esc(b['text'])+' '+' · '.join(hrefs)+'</p>')
        elif typ in ('bullet','list_item'):
            doc.add_paragraph(b['text'],'List Bullet');content.append('<ul><li>'+esc(b['text'])+'</li></ul>')
        else:raise ValueError('Unsupported block '+typ)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();r.font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    doc.core_properties.title=m['title'];doc.core_properties.subject=m.get('subtitle','');doc.core_properties.author='';doc.core_properties.comments=''
    doc.core_properties.keywords='LIBERA, 생성형 AI, 산업배분, 공개 투자판단, 실험 결과'
    doc.save(DOCX)
    toc='<nav aria-label="목차"><strong>목차</strong><ol>'+''.join('<li><a href="#'+h['id']+'">'+esc(h['text'])+'</a></li>' for h in headings if h['level']==1)+'</ol></nav>'
    style='body{margin:0;background:#f3f4f5;color:#141414;font-family:"Malgun Gothic","Noto Sans KR",Arial,sans-serif;font-size:16px;line-height:1.85}main{max-width:1000px;margin:32px auto;padding:56px 64px;background:white}h1,h2,h3,h4{color:#000;line-height:1.45;word-break:keep-all}h1{font-size:30px;margin:0 0 10px}h2{font-size:24px;margin:44px 0 16px}h3{font-size:19px;margin:32px 0 12px}p{margin:0 0 16px;overflow-wrap:anywhere}.subtitle{font-size:19px;margin:0 0 8px}.meta,.note{font-size:13px;color:#444}nav{margin:36px 0;padding:0}nav ol{padding-left:22px}a{color:#174a72;text-underline-offset:3px}figure{margin:24px 0}figcaption{font-weight:700;margin-bottom:8px}.table-wrap{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.6}th,td{border:1px solid #d9d9d9;padding:10px 12px;vertical-align:middle}th{background:#e8eef3;color:black}td{overflow-wrap:anywhere}.note{margin-top:8px}.reference{font-size:14px;padding-left:20px;text-indent:-20px} @media(max-width:800px){main{margin:0;padding:28px 20px}h1{font-size:25px}table{min-width:620px}} @media print{body{background:white;font-size:10pt}main{margin:0;padding:0;max-width:none}nav{display:none}h2,h3,h4,figcaption{break-after:avoid}tr{break-inside:avoid}thead{display:table-header-group}@page{size:A4;margin:20mm 25mm}}'
    style+='nav ol{list-style:none;padding-left:0}th,td{white-space:pre-line}'
    HTML.write_text('<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+esc(m['title'])+'</title><style>'+style+'</style></head><body><main>'+content[0]+toc+''.join(content[1:])+'</main></body></html>',encoding='utf-8')
    QA.mkdir(parents=True,exist_ok=True)
    receipt={'source':str(SOURCE),'source_sha256':digest(SOURCE),'docx':str(DOCX),'docx_sha256':digest(DOCX),'html':str(HTML),'html_sha256':digest(HTML),'blocks':len(m['blocks']),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'headings':headings,'runtime_python':sys.executable,'style_authority':'Earlier v493 builder formal A4 Malgun Gothic style; revised paper content and tables','frozen_experiment_files_modified':False}
    (QA/'build_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in receipt.items() if k!='headings'},ensure_ascii=False))

if __name__=='__main__':main()
