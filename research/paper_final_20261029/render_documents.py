"""Render the two final Korean Markdown documents without changing their prose.

Run with Python 3.13 from the workspace root. Positional paths override defaults.
Images must be local; use readable Unicode, not LaTeX, inside math fences.
"""
from __future__ import annotations

import argparse
import base64
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
from html import escape
from html.parser import HTMLParser
import json
import math
import mimetypes
from pathlib import Path
import re
import subprocess
import sys
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "tmp/paper_document_dependencies"))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from markdown_it import MarkdownIt
from PIL import Image

VERSION = "libera_markdown_renderer_v1"
DEFAULTS = [
    ROOT / "paper/LIBERA_소논문_최종결과반영_20261029.md",
    ROOT / "paper/LIBERA_본실험_결과해석_20261029.md",
]
FONT = "Malgun Gothic"
WIDTH_MM = 166.0
MAX_IMAGE_HEIGHT_MM = 205.0

CSS = """
:root{color-scheme:light;--ink:#1b2633;--muted:#596777;--line:#cdd6df;--accent:#1b5277}
*{box-sizing:border-box}body{margin:0;background:#eef1f4;color:var(--ink);font:16px/1.85 'Malgun Gothic','Apple SD Gothic Neo',sans-serif;word-break:keep-all;overflow-wrap:anywhere}
.sheet{max-width:1020px;margin:32px auto;padding:48px 64px 60px;background:white;border:1px solid #d7dee5;box-shadow:0 3px 20px #162c4410}
nav{margin:0 0 34px;padding:16px 21px;background:#f5f7f9;border-left:3px solid #6a8ba5;font-size:14px;line-height:1.65}nav summary{font-weight:700;cursor:pointer}nav ol{margin:10px 0 0;padding-left:22px}nav li{margin:3px 0}.toc-sub{margin-left:18px;font-size:13px}
h1{margin:0 0 22px;font-size:29px;line-height:1.5;letter-spacing:-.6px;border-bottom:2px solid #446b87;padding-bottom:18px}h2{font-size:22px;line-height:1.5;margin:38px 0 16px;color:#173e58}h3{font-size:18px;line-height:1.5;margin:27px 0 12px}h4{font-size:16px;margin:22px 0 10px}p{margin:0 0 14px}a{color:#175c8a;text-decoration:none;border-bottom:1px solid #b0c6d5}a:hover{text-decoration:underline}strong{font-weight:700}em{font-style:italic}
.table-wrap{overflow-x:auto;margin:20px 0 24px}table{border-collapse:collapse;width:100%;font-size:13.5px;line-height:1.65;font-variant-numeric:tabular-nums}th,td{border:1px solid var(--line);padding:8px 10px;text-align:left;vertical-align:top}th{background:#eaf0f5;font-weight:700}tbody tr:nth-child(even){background:#fafbfc}
img{display:block;max-width:100%;height:auto;margin:22px auto 12px}p:has(>img:only-child){margin-bottom:0}.caption{font-size:13px;color:#526170;line-height:1.65;margin:6px 0 22px}blockquote{border-left:3px solid #b5c5d2;margin:18px 0;padding:4px 20px;color:#415568;background:#f7f9fb}blockquote p:last-child{margin-bottom:0}ul,ol{padding-left:25px;margin:10px 0 18px}li{margin:5px 0}
pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f5f7fa;border:1px solid #dde4eb;padding:14px 18px;font:13px/1.65 Consolas,'Malgun Gothic',monospace}code{font-family:Consolas,'Malgun Gothic',monospace;font-size:.9em}.math{font-family:'Cambria Math','Malgun Gothic',serif;font-size:16px;white-space:pre-wrap;text-align:center;background:#f8fafc;border:0;padding:13px 16px;margin:20px 0}hr{border:0;border-top:1px solid var(--line);margin:28px 0}.render-note{font-size:11px;color:#6b7580;margin-top:38px;padding-top:12px;border-top:1px solid #e2e6eb}
@media(max-width:720px){.sheet{margin:0;padding:28px 20px;border:0;box-shadow:none}body{font-size:15px;line-height:1.8}h1{font-size:25px}h2{font-size:21px}nav{padding:12px 15px}table{font-size:12px}th,td{padding:7px}}
@media print{@page{size:A4;margin:22mm}body{background:white;font-size:10.5pt;line-height:1.5}.sheet{max-width:none;margin:0;padding:0;border:0;box-shadow:none}nav,.render-note{display:none}h1{font-size:19pt}h2{font-size:14pt;margin-top:18pt}h3{font-size:12pt}h1,h2,h3,h4{break-after:avoid}p{orphans:2;widows:2}table{font-size:8.5pt;line-height:1.4}thead{display:table-header-group}tr{break-inside:avoid}.table-wrap{overflow:visible}img{max-height:205mm;object-fit:contain;break-inside:avoid}p:has(>img:only-child){break-after:avoid}.caption{font-size:9pt;break-after:auto;break-inside:avoid}a{color:inherit;border:0;text-decoration:none}pre,.math{font-size:10pt;break-inside:avoid}}
"""


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def plain(children):
    result = []
    for token in children or []:
        if token.type in {"text", "code_inline"}:
            result.append(token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            result.append("\n")
        elif token.type == "html_inline":
            raise ValueError("Raw HTML is not supported in the Markdown source.")
        elif token.type == "image":
            continue
    return "".join(result)


def normalized(text):
    return re.sub(r"\s+", "", text)


def resolve_local(source, url):
    parts = urlsplit(url)
    if parts.scheme or parts.netloc:
        return None
    if not parts.path:
        return None
    return (source.parent / unquote(parts.path)).resolve()


def font_run(run, size=None, bold=False, italic=False, code=False):
    # Unmarked text inherits its paragraph style (including bold headings).
    run.bold = True if bold else None
    run.italic = True if italic else None
    run.font.name = "Consolas" if code else FONT
    if size is not None:
        run.font.size = Pt(size)
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), FONT)
    return run


def style_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def setup_document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(22)
    section.header_distance = section.footer_distance = Mm(10)
    normal = doc.styles["Normal"]
    style_font(normal, 10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [("Title", 19, 0, 16), ("Heading 1", 14, 15, 7), ("Heading 2", 12, 11, 5), ("Heading 3", 11, 8, 4), ("Heading 4", 10.5, 7, 4)]:
        st = doc.styles[name]
        style_font(st, size, bold=True)
        st.font.color.rgb = RGBColor.from_string("173E58" if name != "Title" else "172B3B")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True
        st.paragraph_format.line_spacing = 1.25
    style_font(doc.styles["Caption"], 9)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string("526170")
    doc.styles["Caption"].paragraph_format.line_spacing = 1.35
    doc.styles["Caption"].paragraph_format.space_after = Pt(9)
    for name in ("List Bullet", "List Number"):
        style_font(doc.styles[name], 10.5)
    for name, font, size in [("LIBERA Equation", "Cambria Math", 11), ("LIBERA Code", "Consolas", 9)]:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style_font(st, size)
        st.font.name = font
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_together = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(footer.add_run(), 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    field.append(r)
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.subject = "LIBERA 실험 결과 및 해석"
    doc.core_properties.author = ""
    doc.core_properties.comments = "Rendered from the accompanying Markdown source without prose editing."
    return doc


def xml_text(element):
    chunks = []
    for child in element.iter():
        if child.tag in {qn("w:t"), qn("m:t")}:
            chunks.append(child.text or "")
        elif child.tag == qn("w:br"):
            chunks.append("\n")
        elif child.tag == qn("w:tab"):
            chunks.append("\t")
    return "".join(chunks)


class Renderer:
    def __init__(self, source, tokens):
        self.source, self.tokens = source, tokens
        self.doc = setup_document(next((plain(t.children) for i, t in enumerate(tokens) if t.type == "inline" and i and tokens[i-1].type == "heading_open"), source.stem))
        self.assets, self.links, self.blocks = [], [], []
        self.image_count = self.math_lines = 0
        self.heading_count = self.table_count = 0
        self.bookmarks = 0

    def record(self, paragraph, expected):
        if expected:
            self.blocks.append({"expected": expected, "actual": xml_text(paragraph._p)})

    def hyperlink(self, paragraph, url, children, size, bold, italic):
        link = OxmlElement("w:hyperlink")
        if url.startswith("#"):
            link.set(qn("w:anchor"), url[1:])
        else:
            link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        paragraph._p.append(link)
        local = resolve_local(self.source, url)
        self.links.append({"url": url, "local_path": str(local) if local else None, "exists": local.exists() if local else None})
        # Use regular runs so all Korean and emphasis properties remain editable.
        before = len(paragraph.runs)
        self.inline(paragraph, children, size=size, bold=bold, italic=italic, within_link=True)
        runs = paragraph.runs[before:]
        for run in runs:
            run.font.color.rgb = RGBColor.from_string("175C8A")
            run.underline = True
            link.append(run._element)

    def picture(self, paragraph, token, width_mm=WIDTH_MM):
        url = token.attrGet("src")
        local = resolve_local(self.source, url)
        if local is None or not local.is_file():
            raise FileNotFoundError(f"A local image is required: {url}")
        docx_asset = local
        if local.suffix.lower() == ".svg":
            docx_asset = local.with_suffix(".png")
            if not docx_asset.is_file():
                raise ValueError(f"DOCX needs a matching PNG for SVG: {local}")
        with Image.open(docx_asset) as im:
            width, height = im.size
        draw_width = min(width_mm, MAX_IMAGE_HEIGHT_MM * width / height)
        picture = paragraph.add_run().add_picture(str(docx_asset), width=Mm(draw_width))
        picture._inline.docPr.set("descr", token.content)
        picture._inline.docPr.set("title", token.content)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        self.image_count += 1
        self.assets.append({"source": url, "path": str(local), "sha256": digest(local), "docx_asset": str(docx_asset), "docx_asset_sha256": digest(docx_asset), "width_px": width, "height_px": height, "docx_width_mm": draw_width, "docx_height_mm": draw_width * height / width})

    def inline(self, paragraph, children, *, size=None, bold=False, italic=False, within_link=False, image_width=WIDTH_MM):
        children = children or []
        i = 0
        strong = int(bold)
        emph = int(italic)
        while i < len(children):
            token = children[i]
            if token.type == "strong_open": strong += 1
            elif token.type == "strong_close": strong -= 1
            elif token.type == "em_open": emph += 1
            elif token.type == "em_close": emph -= 1
            elif token.type in {"text", "code_inline"}:
                font_run(paragraph.add_run(token.content), size, bool(strong), bool(emph), token.type == "code_inline")
            elif token.type in {"softbreak", "hardbreak"}:
                paragraph.add_run().add_break()
            elif token.type == "image":
                self.picture(paragraph, token, image_width)
            elif token.type == "link_open":
                end = i + 1
                while end < len(children) and children[end].type != "link_close": end += 1
                if end == len(children): raise ValueError("Unclosed Markdown link")
                self.hyperlink(paragraph, token.attrGet("href"), children[i+1:end], size, bool(strong), bool(emph))
                i = end
            elif token.type not in {"link_close"}:
                raise ValueError(f"Unsupported inline Markdown token: {token.type}")
            i += 1

    def paragraph(self, token, *, style=None, list_prefix=None, indent=0):
        text = plain(token.children)
        if style is None and re.match(r"^(?:그림|표)\s*[A-Z]?\d+(?:[.．:\s]|$)", text):
            style = "Caption"
        p = self.doc.add_paragraph(style=style)
        if list_prefix:
            font_run(p.add_run(list_prefix), 10.5)
            p.paragraph_format.left_indent = Mm(5 + indent * 5)
            p.paragraph_format.first_line_indent = Mm(-4)
        self.inline(p, token.children)
        if style == "Caption":
            # A figure stays with its caption, but the caption does not pull
            # the following body paragraph onto the same page.
            p.paragraph_format.keep_with_next = text.startswith("표")
            p.paragraph_format.keep_together = True
        self.record(p, (list_prefix or "") + text)
        return p

    def table(self, tokens):
        rows, row, cell = [], None, None
        for token in tokens:
            if token.type == "tr_open": row = []
            elif token.type in {"th_open", "td_open"}: cell = []
            elif token.type == "inline": cell.extend(token.children or [])
            elif token.type in {"th_close", "td_close"}: row.append(cell)
            elif token.type == "tr_close": rows.append(row)
        if not rows or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError("Malformed Markdown table")
        ncols = len(rows[0])
        sizes = []
        for col in range(ncols):
            lengths = [sum(2 if ord(c) > 127 else 1 for c in plain(row[col])) for row in rows]
            sizes.append(math.sqrt(max(6, min(85, sorted(lengths)[min(len(lengths)-1, int(.8*len(lengths)))]))))
        widths = [WIDTH_MM * value / sum(sizes) for value in sizes]
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        props = table._tbl.tblPr
        total = props.find(qn("w:tblW"))
        total.set(qn("w:type"), "dxa")
        total.set(qn("w:w"), str(round(Mm(WIDTH_MM).twips)))
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement("w:" + side)
            for key, value in {"val": "single", "sz": "4", "color": "CDD6DF"}.items(): edge.set(qn("w:"+key), value)
            borders.append(edge)
        props.append(borders)
        for j, width in enumerate(widths): table.columns[j].width = Mm(width)
        for r, row in enumerate(rows):
            trpr = table.rows[r]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if r == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for c, content in enumerate(row):
                target = table.cell(r, c)
                target.width = Mm(widths[c])
                target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = target._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:"+side)
                    margin.set(qn("w:w"), "75")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                tcpr.append(margins)
                if r == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EAF0F5")
                    tcpr.append(shade)
                p = target.paragraphs[0]
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                # Short comparison tables should be read as one unit. Long
                # reproducibility tables can continue with repeated headers.
                p.paragraph_format.keep_with_next = (r < len(rows) - 1) if len(rows) <= 10 else r == 0
                text = plain(content)
                if c > 0 and (r == 0 or re.fullmatch(r"[\d\s.,%+−\-–/:()≤≥=]+", text)):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.inline(p, content, size=8.5 if ncols >= 7 else 9, bold=r == 0, image_width=widths[c]-3)
                self.record(p, text)
        gap = self.doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(3)
        gap.paragraph_format.space_before = Pt(0)
        gap.paragraph_format.line_spacing = .4
        self.table_count += 1

    def fence(self, token):
        math_block = token.info.strip().lower() in {"math", "equation", "equations"}
        if math_block and re.search(r"\\[A-Za-z]{2,}|\$\$", token.content):
            raise ValueError("Math fences must use readable Unicode, not LaTeX commands.")
        for line in token.content.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph(style="LIBERA Equation" if math_block else "LIBERA Code")
            if math_block:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                equation = OxmlElement("m:oMath")
                run = OxmlElement("m:r")
                text = OxmlElement("m:t")
                text.set(qn("xml:space"), "preserve")
                text.text = line
                run.append(text)
                equation.append(run)
                p._p.append(equation)
                self.math_lines += 1
            else:
                font_run(p.add_run(line), 9, code=True)
            self.record(p, line)

    def build(self):
        i, lists, blockquote = 0, [], 0
        item_first = []
        while i < len(self.tokens):
            t = self.tokens[i]
            if t.type == "heading_open":
                inline = self.tokens[i+1]
                level = int(t.tag[1])
                p = self.paragraph(inline, style="Title" if level == 1 else "Heading " + str(min(level-1, 4)))
                if level == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.bookmarks += 1
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), str(self.bookmarks))
                start.set(qn("w:name"), t.attrGet("id"))
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), str(self.bookmarks))
                # Paragraph properties must precede bookmark/run children in OOXML.
                p._p.insert(1 if p._p.pPr is not None else 0, start)
                p._p.append(end)
                self.heading_count += 1
                i += 2
            elif t.type == "paragraph_open":
                inline = self.tokens[i+1]
                prefix = None
                if lists and item_first and item_first[-1]:
                    if lists[-1]["kind"] == "ordered":
                        prefix = str(lists[-1]["next"]) + ". "
                        lists[-1]["next"] += 1
                    else: prefix = "• "
                    item_first[-1] = False
                p = self.paragraph(inline, list_prefix=prefix, indent=max(0, len(lists)-1))
                if blockquote:
                    p.paragraph_format.left_indent = Mm(7)
                    p.paragraph_format.right_indent = Mm(5)
                i += 2
            elif t.type in {"bullet_list_open", "ordered_list_open"}:
                lists.append({"kind": "ordered" if t.type.startswith("ordered") else "bullet", "next": int(t.attrGet("start") or 1)})
            elif t.type in {"bullet_list_close", "ordered_list_close"}: lists.pop()
            elif t.type == "list_item_open": item_first.append(True)
            elif t.type == "list_item_close": item_first.pop()
            elif t.type == "blockquote_open": blockquote += 1
            elif t.type == "blockquote_close": blockquote -= 1
            elif t.type == "table_open":
                end = i + 1
                while self.tokens[end].type != "table_close": end += 1
                self.table(self.tokens[i+1:end])
                i = end
            elif t.type in {"fence", "code_block"}: self.fence(t)
            elif t.type == "hr":
                p = self.doc.add_paragraph()
                border = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for key, value in {"val":"single", "sz":"4", "color":"CDD6DF"}.items(): bottom.set(qn("w:"+key), value)
                border.append(bottom)
                p._p.get_or_add_pPr().append(border)
            elif t.type not in {"paragraph_close", "heading_close"}:
                raise ValueError(f"Unsupported block Markdown token: {t.type}")
            i += 1
        return self.doc


def prepare_tokens(md, text):
    tokens = md.parse(text)
    headings = []
    for i, token in enumerate(tokens):
        if token.type == "heading_open":
            anchor = "section-" + str(len(headings)+1)
            token.attrSet("id", anchor)
            headings.append({"id": anchor, "level": int(token.tag[1]), "text": plain(tokens[i+1].children)})
        elif token.type == "paragraph_open" and i+1 < len(tokens):
            if re.match(r"^(?:그림|표)\s*[A-Z]?\d+(?:[.．:\s]|$)", plain(tokens[i+1].children)):
                token.attrSet("class", "caption")
    return tokens, headings


def html_document(md, tokens, headings, source):
    tokens = deepcopy(tokens)
    embedded = []
    for parent in tokens:
        for child in parent.children or []:
            if child.type == "image":
                path = resolve_local(source, child.attrGet("src"))
                if path is None or not path.is_file(): raise FileNotFoundError(child.attrGet("src"))
                mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
                child.attrSet("src", f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode("ascii"))
                embedded.append(str(path))
    def fence(tokens_, index, options, env):
        token = tokens_[index]
        cls = "math" if token.info.strip().lower() in {"math", "equation", "equations"} else "code"
        return '<pre class="' + cls + '">' + escape(token.content) + '</pre>\n'
    md.renderer.rules["fence"] = fence
    md.renderer.rules["table_open"] = lambda *args: '<div class="table-wrap"><table>\n'
    md.renderer.rules["table_close"] = lambda *args: '</table></div>\n'
    title = headings[0]["text"] if headings else source.stem
    toc = "".join('<li' + (' class="toc-sub"' if h["level"] > 2 else '') + '><a href="#' + h["id"] + '">' + escape(h["text"]) + '</a></li>' for h in headings if h["level"] in {2, 3})
    body = md.renderer.render(tokens, md.options, {})
    document = '<!doctype html>\n<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + escape(title) + '</title><style>' + CSS + '</style></head><body><main class="sheet">'
    if toc: document += '<nav aria-label="목차"><details open><summary>목차</summary><ol>' + toc + '</ol></details></nav>'
    document += '<article id="document-content">' + body + '</article></main></body></html>\n'
    return document, embedded


class ArticleText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.inside = False
        self.text = []
    def handle_starttag(self, tag, attrs):
        if tag == "article": self.inside = True
    def handle_endtag(self, tag):
        if tag == "article": self.inside = False
    def handle_data(self, text):
        if self.inside: self.text.append(text)


def export_pdf(docx, qa_dir):
    pdf = docx.with_suffix(".pdf")
    receipt = qa_dir / (docx.stem + ".word_export.json")
    command = ["powershell.exe", "-NoProfile", "-File", str(ROOT / "research/solution_revision_20261026/render_word.ps1"),
               "-InputPath", str(docx), "-PdfPath", str(pdf), "-ReceiptPath", str(receipt)]
    result = subprocess.run(command, cwd=ROOT, text=True, encoding="utf-8", errors="replace", capture_output=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if result.returncode:
        raise RuntimeError("Word PDF export failed: " + result.stdout + result.stderr)
    return {"pdf": str(pdf), "receipt": str(receipt), "sha256": digest(pdf), "receipt_sha256": digest(receipt)}


def same_path(recorded, expected):
    return isinstance(recorded, str) and Path(recorded).resolve() == expected.resolve()


def owned_output_plan(source, qa_dir, replace_rendered):
    """Validate every existing output before planning a nonrecursive archive move."""
    html_path, docx_path, pdf_path = (source.with_suffix(ext) for ext in (".html", ".docx", ".pdf"))
    qa_path = qa_dir / (source.stem + ".render_qa.json")
    word_path = qa_dir / (source.stem + ".word_export.json")
    candidates = [html_path, docx_path, qa_path, pdf_path, word_path]
    if not any(path.exists() for path in candidates):
        return []
    if not replace_rendered or not qa_path.is_file():
        raise FileExistsError("Outputs exist. --replace-rendered requires a matching prior renderer QA receipt.")
    previous = json.loads(qa_path.read_text(encoding="utf-8-sig"))
    if (previous.get("renderer") != VERSION or previous.get("status") != "passed"
            or not same_path(previous.get("source_path"), source)):
        raise ValueError("Existing outputs do not have a matching successful renderer receipt.")

    expected_hashes = {}
    for label, path in (("html", html_path), ("docx", docx_path)):
        if not same_path(previous.get(label + "_path"), path) or not path.is_file():
            raise ValueError(f"Prior {label} output path is missing or does not match the receipt: {path}")
        expected_hashes[path] = previous.get(label + "_sha256")

    prior_pdf = previous.get("pdf_export")
    if prior_pdf or pdf_path.exists() or word_path.exists():
        if (not isinstance(prior_pdf, dict) or not pdf_path.is_file() or not word_path.is_file()
                or not same_path(prior_pdf.get("pdf"), pdf_path)
                or not same_path(prior_pdf.get("receipt"), word_path)):
            raise ValueError("Existing PDF and Word receipt are not linked by the prior renderer QA; nothing was moved.")
        word = json.loads(word_path.read_text(encoding="utf-8-sig"))
        if (word.get("status") != "completed"
                or word.get("renderer") != "native_Microsoft_Word_independent_COM_fallback"
                or not same_path(word.get("input_docx"), docx_path)
                or not same_path(word.get("pdf_path"), pdf_path)
                or word.get("input_sha256_before") != expected_hashes[docx_path]
                or word.get("input_sha256_after") != expected_hashes[docx_path]
                or word.get("pdf_sha256") != prior_pdf.get("sha256")):
            raise ValueError("Existing Word receipt does not match the prior DOCX/PDF hashes; nothing was moved.")
        expected_hashes[pdf_path] = prior_pdf.get("sha256")
        # Older renderer receipts did not store this hash. Their Word receipt
        # is authenticated by the path, renderer, status, and both file hashes.
        expected_hashes[word_path] = prior_pdf.get("receipt_sha256") or digest(word_path)
    expected_hashes[qa_path] = digest(qa_path)

    plan = []
    workspace = ROOT.resolve()
    if not qa_dir.resolve().is_relative_to(workspace):
        raise ValueError("Archive directory must remain inside the workspace.")
    for path, expected_hash in expected_hashes.items():
        resolved = path.resolve()
        if path.is_symlink() or not resolved.is_relative_to(workspace) or resolved == source:
            raise ValueError(f"Refusing an output path outside the workspace or aliased to the source: {path}")
        if not isinstance(expected_hash, str) or digest(path) != expected_hash:
            raise ValueError(f"Existing output was changed after rendering; nothing was moved: {path}")
        plan.append({"path": str(resolved), "sha256": expected_hash})
    return plan


def archive_owned_outputs(plan, qa_dir):
    """Move only individually verified files; never delete or recursively move."""
    if not plan:
        return None
    # Recheck immediately before mutation in case a file changed while the new
    # source was parsed and validated. All new content is ready at this point.
    for item in plan:
        path = Path(item["path"])
        if not path.is_file() or path.is_symlink() or digest(path) != item["sha256"]:
            raise ValueError(f"Output changed before archival; nothing was moved: {path}")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    archive = (qa_dir / ("prior_render_" + stamp)).resolve()
    if not archive.is_relative_to(ROOT.resolve()):
        raise ValueError("Archive target must remain inside the workspace.")
    archive.mkdir(parents=True, exist_ok=False)
    manifest_path = archive / "archive_manifest.json"
    manifest = {"schema": "libera_prior_render_archive_v1", "renderer": VERSION,
                "status": "planned", "files": []}
    for item in plan:
        destination = archive / Path(item["path"]).name
        if destination.exists() or destination.parent != archive:
            raise ValueError(f"Archive destination is not a new direct child: {destination}")
        manifest["files"].append({"original_path": item["path"], "archived_path": str(destination),
                                  "sha256": item["sha256"], "moved": False})
    def save_manifest():
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    save_manifest()
    try:
        for item in manifest["files"]:
            Path(item["original_path"]).rename(item["archived_path"])
            item["moved"] = True
            if digest(Path(item["archived_path"])) != item["sha256"]:
                raise ValueError("Archived file hash does not match its original receipt.")
            save_manifest()
        manifest["status"] = "completed"
        save_manifest()
    except Exception as error:
        manifest["status"] = "partial_archive_preserved"
        manifest["error"] = str(error)
        save_manifest()
        raise RuntimeError(f"Archival stopped; all files were preserved. Inspect {manifest_path}") from error
    return {"directory": str(archive), "manifest": str(manifest_path), "files_preserved": len(plan)}


def render(source, qa_dir, replace_rendered=False, pdf=False):
    source = source.resolve()
    if source.suffix.lower() != ".md" or not source.is_file(): raise FileNotFoundError(source)
    html_path, docx_path = source.with_suffix(".html"), source.with_suffix(".docx")
    qa_path = qa_dir / (source.stem + ".render_qa.json")
    archive_plan = owned_output_plan(source, qa_dir, replace_rendered)
    source_hash = digest(source)
    md = MarkdownIt("commonmark", {"html": False, "breaks": False}).enable("table")
    tokens, headings = prepare_tokens(md, source.read_text(encoding="utf-8-sig"))
    renderer = Renderer(source, tokens)
    doc = renderer.build()
    bad_blocks = [{"index": i, **block} for i, block in enumerate(renderer.blocks) if block["expected"] != block["actual"]]
    missing_links = [x for x in renderer.links if x["local_path"] and not x["exists"]]
    if bad_blocks: raise ValueError("DOCX text differs from source: " + repr(bad_blocks[:3]))
    if missing_links: raise FileNotFoundError("Local links are missing: " + repr(missing_links))
    html, embedded = html_document(md, tokens, headings, source)
    parser = ArticleText()
    parser.feed(html)
    html_text = normalized("".join(parser.text))
    # Lists add display prefixes to DOCX; the Markdown list marker is not text in HTML.
    missing_html = []
    for block in renderer.blocks:
        expected = block["expected"]
        stripped = re.sub(r"^(?:• |\d+\. )", "", expected)
        if normalized(stripped) not in html_text: missing_html.append(expected)
    if missing_html: raise ValueError("HTML text is missing source content: " + repr(missing_html[:3]))
    qa_dir.mkdir(parents=True, exist_ok=True)
    prior_archive = archive_owned_outputs(archive_plan, qa_dir)
    doc.save(docx_path)
    html_path.write_text(html, encoding="utf-8", newline="\n")
    with ZipFile(docx_path) as package:
        tree = etree.fromstring(package.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "m":"http://schemas.openxmlformats.org/officeDocument/2006/math"}
        table_properties = [{"header_repeats": bool(t.xpath("./w:tr[1]/w:trPr/w:tblHeader", namespaces=ns)),
                             "all_rows_no_split": len(t.xpath("./w:tr/w:trPr/w:cantSplit", namespaces=ns)) == len(t.xpath("./w:tr", namespaces=ns))}
                            for t in tree.xpath("//w:tbl", namespaces=ns)]
        body_paragraphs = len(tree.xpath("/w:document/w:body/w:p", namespaces=ns))
        all_paragraphs = len(tree.xpath("//w:p", namespaces=ns))
        hyperlink_count = len(tree.xpath("//w:hyperlink", namespaces=ns))
        equation_count = len(tree.xpath("//m:oMath", namespaces=ns))
    unchanged = digest(source) == source_hash
    qa = {
        "schema": "libera_document_render_qa_v1", "renderer": VERSION,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "passed" if unchanged and all(all(t.values()) for t in table_properties) else "failed",
        "source_path": str(source), "source_sha256": source_hash, "source_unchanged": unchanged,
        "html_path": str(html_path), "html_sha256": digest(html_path), "docx_path": str(docx_path), "docx_sha256": digest(docx_path),
        "counts": {"headings": renderer.heading_count, "source_text_blocks_checked": len(renderer.blocks), "docx_body_paragraphs": body_paragraphs,
                   "docx_all_paragraphs": all_paragraphs, "tables": renderer.table_count, "images": renderer.image_count,
                   "docx_hyperlinks": hyperlink_count, "editable_math_lines": equation_count, "html_embedded_images": len(embedded)},
        "text_checks": {"docx_mismatches": bad_blocks, "html_missing_blocks": missing_html},
        "table_checks": table_properties, "assets": renderer.assets, "links": renderer.links, "missing_assets": [], "missing_local_links": missing_links,
        "html_self_contained_images": True, "layout": {"paper": "A4", "margins_mm": 22, "font": FONT, "body_pt": 10.5, "line_spacing": 1.45,
                                                        "body_height_mm": 253, "max_image_height_mm": MAX_IMAGE_HEIGHT_MM,
                                                        "image_keep_with_next": True, "figure_caption_keep_with_next": False},
        "prior_render_archive": prior_archive,
        "visual_review_required": True,
    }
    qa_path.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if pdf:
        try:
            qa["pdf_export"] = export_pdf(docx_path, qa_dir)
        except Exception as error:
            qa["pdf_export_error"] = str(error)
            qa_path.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            raise
        qa_path.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"source": str(source), "status": qa["status"], "counts": qa["counts"], "qa": str(qa_path)}, ensure_ascii=False))
    if qa["status"] != "passed": raise RuntimeError("Document QA failed")
    return qa


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="*", type=Path)
    parser.add_argument("--qa-dir", type=Path, default=HERE / "document_qa")
    parser.add_argument("--replace-rendered", action="store_true", help="Archive only outputs whose paths and hashes match the prior renderer/Word receipts, then regenerate.")
    parser.add_argument("--pdf", action="store_true", help="Export via the existing isolated Word wrapper; verified prior exports are preserved with --replace-rendered.")
    args = parser.parse_args()
    sources = args.paths or DEFAULTS
    for source in sources:
        render(source, args.qa_dir.resolve(), args.replace_rendered, args.pdf)


if __name__ == "__main__":
    main()
