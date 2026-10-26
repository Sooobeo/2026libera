"""Restore the supplied abstract as the study concept and retain completed pilot data."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import argparse
import importlib.util
import json
import re

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
BASE = ROOT / 'paper/v530_manuscript_20261026.json'
OUT = ROOT / 'paper/v540_manuscript_20261026.json'
STEM = 'LIBERA_소논문_원초록반영_v5.4_20261026'
TITLE = '공통 GenAI는 기관투자자의 전략적 이질성을 축소하는가?'
SUBTITLE = '-액티브 주식형 펀드의 자산배분 프로세스 내 Workflow RAG 배치와 완화효과를 중심으로-'
PDF_P1 = '본 연구는 공통 생성형 AI(GenAI)의 활용이 기관투자자의 전략적 이질성과 투자판단 품질에 미치는 영향을 검토하고, 기관별 운용조건에 적합한 판단을 유지하기 위한 대응방안을 모색한다. Morgan Stanley는 2024년 자체 리서치의 검색과 요약을 지원하는 생성형 AI 도구 AskResearchGPT를 출시했다(Morgan Stanley, 2024). 영국은행과 FCA의 2024년 조사에서도 응답 금융회사의 75%가 AI를 사용하고 있었으며, 전체 AI 활용사례 중 완전자율 의사결정의 비중은 2%였다(Bank of England & FCA, 2024). 이러한 배경에서 본 연구는 AI의 최종 결정뿐 아니라, 투자판단에 필요한 정보를 처리하고 분석하는 중간 업무에서 AI의 영향이 어떻게 누적되는지에 주목한다.'
PDF_P2 = '액티브 자산운용에서는 동일한 시장정보를 접하더라도 펀드의 투자목표, 기존 보유자산과 위험제약에 따라 서로 다른 배분 판단이 요구된다. 이러한 차이는 투자자의 요구에 맞는 운용과 기관별 전략의 차별화를 뒷받침한다. 그러나 자료 검색과 요약, 전망 분석, 배분안 작성에 유사한 GenAI가 분산적으로 사용되면, 각 단계의 산출물이 다음 단계의 정보 선택과 판단 기준에 영향을 미칠 수 있다. 이에 서로 다른 기관이 명시적으로 동일한 투자 규칙을 채택하지 않더라도 공통모델의 해석 방식에 의존하면서, 기관별로 달라야 할 판단까지 유사해질 가능성이 있다. 본 연구는 개별 조건상 필요한 배분 차이가 기준보다 작아지는 현상을 ‘과잉수렴’으로 정의하고, 그것이 판단 손실이나 투자제약 위반을 동반하는지 함께 검토한다. 이는 개별 펀드의 운용 적합성뿐 아니라, 유사한 투자행동과 위험 노출의 확대 가능성에도 연결되는 금융적 문제다.'

def p(text): return {'type': 'paragraph', 'text': text}
def h(text, level=1): return {'type': 'heading', 'text': text, 'level': level}
def digest(path): return sha256(path.read_bytes()).hexdigest()
def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    obj = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(obj)
    return obj

def reframe_as_main_experiment(m, previous):
    """Apply the user's explicit positioning: completed five-condition study is the main experiment."""
    blocks = m['blocks']
    def section(title, body):
        start = next(i for i,b in enumerate(blocks) if b['type']=='heading' and b['text']==title)
        end = next((i for i in range(start+1,len(blocks)) if blocks[i]['type']=='heading'), len(blocks))
        blocks[start+1:end] = body
    section('국문초록', [p(PDF_P1), p(PDF_P2),
        p('본실험은 동일한 시장 자료와 미국 주식 11개 산업의 새 자금 배분 과제에서 Fidelity와 UBS의 공개 산업전망을 직접 제공한 다섯 조건을 비교하였다. 기관 자료의 초기·후기 제공 및 미제공에 따라 배분과 설명이 어떻게 달라지는지 살펴보고, 기관 원문과 선택 이유, 실제 제출 비중을 대조하였다. 네 기관 자료 조건과 대조군의 파생 배분 거리는 4.76~12.98%였으나 정확한 자료 요약 뒤에도 서로 다른 배분에 같은 수치 오류가 남았고, 설명과 실제 비중이 어긋났다. 외부 정규화가 원비중의 증감 방향을 뒤집는 사례도 확인하였다. 이 결과는 배분의 차이와 함께 근거의 전달 및 판단 품질을 검토해야 함을 보여준다.'),
        p('이 관측을 바탕으로 계획, 배분안 구성, 점검과 수정의 3단계에 기관별 근거를 연결하는 Workflow RAG 대응 설계를 제안한다. 핵심은 근거를 검색하여 제공하는 것과 그 근거가 실제 판단에 남도록 업무를 설계하는 것을 함께 다루는 데 있다. 후속 검증에서는 같은 검색 근거를 어느 단계부터 제공하는지 달리하고 이후에는 유지하며, 추가 근거가 없는 조건과 전체 적격 원문을 처음부터 제공하는 조건도 대조한다. 본실험에서 확인한 설명·비중의 불일치와 공동오류를 각 단계의 검토 항목에 반영한다.'),
        p('실제 미국 액티브 주식형 펀드 4개의 보유자료·투자설명서·운용보고서를 사용하고, 허용 정보와 명시 제약에 따른 수리 기준배분을 계산하는 비교는 확장 검증으로 제시한다. 이 비교에서는 필요한 펀드 간 차이, 개별 정보 반응, 모형상 기대수익 손실과 투자제약 위반을 함께 평가하고, 보류한 새 문서에서 검색을 다시 수행하며 검색·실행 비용도 확인한다. 현재 실험이 과잉수렴이나 RAG의 완화효과를 직접 입증한 것은 아니다. 본 연구는 본실험의 관측과 대응 설계를 연결하여, 필요한 판단 차이와 품질을 유지하기 위한 근거 제공 시점, 검토 절차 및 단계 간 정보 전달의 원칙을 제안한다.'),
        p('주제어  생성형 AI, 기관투자자, 전략적 이질성, 과잉수렴, 자산배분, Workflow RAG')])
    start = next(i for i,b in enumerate(blocks) if b['type']=='heading' and b['text']=='1 서론')
    end = next(i for i in range(start+1,len(blocks)) if blocks[i]['type']=='heading')
    intro = deepcopy(blocks[start+1:end][:5])
    intro[3] = p('본 연구는 개별 조건상 필요한 배분 차이가 기준보다 작아지는 현상을 과잉수렴으로 정의하고, 판단 손실과 투자제약 위반을 함께 검토한다. 배분 간 거리가 작다는 사실만으로 필요한 차이의 소실을 판정할 수는 없다. 공통 알고리즘의 개별 성능과 집단 후생, 공유 모델과 결과 동질화를 나누는 선행연구도 이러한 구분을 뒷받침한다(Kleinberg와 Raghavan, 2021; Bommasani 등, 2022). 본고는 기관별 근거가 실제 선택에 이어지는지 대조하여 이 문제를 구체화한다. 유사한 권고가 실제 거래와 위험 노출의 동조로 이어질 가능성은 금융적 함의이지만, 시장 수준의 연결 자체는 이번 실험의 관측 범위 밖이다(FSB, 2024).')
    intro += [p('본실험은 동일한 시장 자료에서 두 기관의 공개 산업전망을 초기 또는 후기에 직접 제공한 네 조건과 미제공 대조군을 비교하였다. 새 자금을 GICS 11개 산업에 배분하게 하고, 기관 자료의 의미와 모델이 밝힌 이유, 실제 비중을 대조하였다. 이는 공통 GenAI에 기관별 자료를 넣었을 때 판단의 차이가 어떤 근거를 갖는지 살펴보는 소논문의 실험 과정이다. 실제 펀드의 보유조건과 조건부 수리 기준배분을 사용하는 비교는 이후의 확장 검증으로 제시한다.'),
        p('결과에서는 기관 자료 조건의 배분이 대조군과 달랐어도 설명과 실제 선택이 어긋났고, 서로 다른 배분에 같은 수치 오류가 남았다. 원비중과 외부 정규화된 배분의 증감 방향이 반대인 경우도 있었다. 이 관측은 기관 자료를 추가하고 결과가 달라졌다는 사실만으로 필요한 판단 차이와 품질이 유지되었다고 볼 수 없음을 보여준다. 본고는 문제 제기, 본실험, 결과 해석과 대응 설계를 하나의 과정으로 연결하고, 확인한 단절을 Workflow RAG의 근거 전달과 검토 절차에 반영한다. 3단계 검색 배치의 완화효과는 이 대응 설계를 평가할 후속 질문으로 남긴다.')]
    blocks[start+1:end] = intro
    section('3 연구질문과 비교 논리', [
        p('본고는 공통 GenAI에 기관별 자료를 제공했을 때 그 차이가 어떤 근거를 거쳐 배분 선택에 이어지는지를 본실험으로 살펴본다. 전략적 이질성과 판단 품질이라는 문제를 자료 해석, 배분 차이, 제공 절차의 세 연구질문으로 구체화한다.'),
        p('RQ1은 기관별 공개 산업전망의 의미가 설명과 실제 비중 사이에서 유지되는가이다. 원문의 방향·범위·이유를 산업행 요약과 선택 이유에 대조하고, 모델이 밝힌 이유가 제출 비중과 일치하는지 확인한다. 기관 전망을 따르지 않았더라도 원문을 정확히 이해하고 별도 판단의 이유를 밝혔으면 그 자체를 오독으로 분류하지 않는다.'),
        p('RQ2는 기관 전망을 제공한 조건의 최종 배분이 공통 시장 자료만 제공한 대조군과 어떻게 다른가이다. 같은 외부 계산을 적용한 배분 사이의 거리를 비교한다. 이 값은 결과의 차이를 보여주며, 필요한 전략적 차이나 과잉수렴의 크기로 해석하지 않는다.'),
        p('RQ3은 기관 자료를 초안부터 제공한 경우와 최종 검토 때 처음 제공한 경우에 근거 해석과 배분 수정이 어떻게 다른가이다. 같은 기관의 초기·후기 조건 및 각 조건의 초안·최종안을 비교한다. 초안·노출 횟수·대화 이력도 함께 달라지므로 순수한 시점 효과를 추정한 것은 아니며, 관측된 차이를 Workflow RAG의 단계 배치와 검토 절차를 설계하는 근거로 사용한다.')])
    section('4 예비실험의 자료와 설계', [p('본실험은 두 기관의 공개 산업전망과 같은 시장 자료를 사용한 다섯 조건으로 구성하였다. 현재 정리된 GICS 최상위 11개 산업에 새 자금을 배분하도록 하고, 초안 D0와 첫 완료 최종안 F0에서 설명과 비중을 관측하였다. 기존 펀드의 보유자산을 조정하는 과제는 실행하지 않았으며, 실제 펀드의 운용조건과 수리 기준배분을 사용하는 비교는 8.6 및 9장의 확장 검증으로 다룬다.')])
    section('5 예비실험 평가방법', [p('본실험에서는 제공 문서와 설명·비중의 대응, 조건별 배분 차이와 수정 양상을 평가한다. 필요한 펀드 간 차이의 수리 기준과 모형상 판단 손실은 이번 실험에서 계산하지 않았으므로, 산출물 차이와 과잉수렴의 판정을 구분한다.')])
    section('6 예비실험 결과', [p('아래 결과는 세 연구질문에 대응한다. 기관 전망을 제공한 네 조건의 최종 배분은 대조군과 달랐다. 그 차이가 기관 근거의 충실한 반영에서 나왔는지 확인하기 위해 원비중의 유효성, 최종 배분의 차이와 설명·숫자의 대응을 함께 살펴본다.')])
    section('6 5 예비 질문에 대한 답과 본실험의 과제', [p('RQ1에서는 기관별 산업행의 등급과 요약이 보존되었으나 이후 해석과 실제 비중의 연결은 일관되지 않았다. RQ2에서는 네 기관 자료 조건의 파생 배분이 C와 모두 달랐다. RQ3에서는 초기·후기 조건의 수정 양상이 달랐지만 어느 절차가 근거를 더 충실히 반영했는지는 확정하지 못하였다. 이 결과는 배분 거리와 함께 근거 반응, 공동오류와 판단 품질을 검토해야 함을 보여준다. 다음 장에서는 이를 Workflow RAG의 근거 전달과 점검·수정 절차에 연결한다.')])
    section('7 예비결과와 Workflow RAG 설계의 연결', [p('본실험에서 기관 자료가 요약에 반영되어도 실제 선택과의 연결은 끊길 수 있었다. 이 관측을 토대로 검색 결과를 제공한 사실과 판단에 사용한 결과를 나누어 확인하고, 계획·배분안 구성·점검과 수정에 어떤 검사를 연결할지 해석한다. 배분 차이를 인위적으로 키우기보다 기관별 근거가 지지하는 판단을 추적할 수 있는 업무 설계가 필요하다.')])
    section('8 Workflow RAG의 배치와 검증 설계', [p('대응의 목표는 공통 시장정보를 함께 반영하면서 펀드별 근거와 운용조건이 지지하는 차이 및 판단 품질을 유지하는 것이다. 본실험의 관측을 바탕으로 계획, 배분안 구성, 점검과 수정으로 이어지는 Workflow RAG를 제안한다. 근거를 검색해 제공하는 것과 그 근거가 다음 단계의 판단 및 확정 비중에 남도록 업무를 설계하는 것을 함께 다룬다.'),
        p('본실험은 기관 전망을 직접 제공한 초안·최종안의 두 단계로 수행하였다. 자료 요약과 선택의 단절, 서로 다른 배분에 남은 공유 오류와 정규화에 의한 증감 방향의 반전을 확인하였으며, 이 결과를 아래 대응 설계의 검사 항목에 반영한다. 검색기를 사용하는 3단계 배치 비교와 완화효과는 후속 검증 대상이다.')])
    section('8 6 실제 펀드 과제로의 적용과 역할 구분', [
        p('본실험의 공개 전망 조건을 실제 운용조건으로 확장하려면 펀드별 투자목적, 기존 보유자산, 허용 위험과 명시 제약을 추가해야 한다. 미국 액티브 주식형 펀드 4개의 공개 보유자료·투자설명서·운용보고서로 이러한 조건을 구성하고, 현재 정리한 GICS 11개 산업 체계에서 목표 비중과 기존 보유 대비 변경분을 비교하는 확장 검증을 제안한다. 네 펀드의 실행과 수리 기준배분 계산은 이번 완료 결과에 포함되지 않는다.'),
        p('운영 단계에서는 자료 준비자가 적격 원문과 버전을, 운용 담당자가 목표·제약·허용 재량을, 검증 담당자가 설명·숫자의 불일치와 미해결 사항을 관리한다. NIST AI RMF 1.0을 참고한 이 책임 구분은 기관 전망, 모델의 추가 해석, 실제 승인한 비중이 하나의 주장으로 섞이는 것을 막기 위한 설계다(Tabassi, 2023). 기관 고유의 전략을 얼마나 보존하는지는 기록의 존재가 아니라 운용조건과 판단 품질을 포함한 후속 비교에서 평가해야 한다.')])
    section('9 1 예비관찰로 확인한 범위와 남은 질문', [
        p('본실험은 한 시장 상황, 두 공개 자료원, 한 소형 양자화 모델과 조건당 한 최종안을 분석하였다. 초안·최종안의 두 단계에 자료를 직접 제공했으며 3단계 검색 RAG, 네 실제 펀드의 보유조건을 반영한 배분은 실행하지 않았다. 초기·후기 조건에는 노출량·초안·대화 이력과 운영 이력도 함께 달랐다. 초기 출력 뒤에 분석 역할과 정규화된 파생 지표를 변경했으므로 수정량을 근거 제공 시점의 인과효과로 해석할 수 없다.'),
        p('실제 보유비중·고객 목적·위험선호와 미래 성과가 없는 이번 과제에서는 필요한 정답 배분, 모형상 판단 손실 또는 과잉수렴을 직접 판정하지 않는다. 사후 의미 검토도 부분 가림의 AI 보조 방식으로 독립적인 사람 평가를 대신하지 않는다. 본실험은 배분 차이와 정확한 요약이 있어도 근거와 선택의 연결이 어긋날 수 있음을 구체적으로 보여주었으며, 이 관측을 토대로 제안한 Workflow RAG의 개선 효과는 별도로 검증해야 한다.')])
    section('10 결론', [
        p('본 연구는 공통 GenAI가 기관투자자의 전략적 이질성을 축소하는가라는 문제를 기관별 자료의 해석과 배분 선택 과정에서 살펴보았다. 동일한 시장 자료와 투자 과제에서 두 기관의 공개 산업전망을 초기·후기에 제공하거나 제공하지 않은 다섯 조건을 본실험으로 구성하고, 원문·설명·비중의 연결을 대조하였다.'),
        p('본실험에서는 기관 자료 조건의 배분이 대조군과 달랐지만 정확한 산업행 요약 뒤에도 설명과 실제 비중이 어긋났으며, 서로 다른 배분에 같은 수치 오류가 남았다. 외부 정규화가 원비중의 증감 방향을 바꾸는 경우도 확인하였다. 이는 자료를 추가하고 배분이 달라졌다는 관찰만으로 기관별로 필요한 판단 차이와 품질이 유지되었다고 볼 수 없음을 보여준다.'),
        p('이에 계획, 배분안 구성, 점검과 수정의 3단계에 기관별 근거를 연결하는 Workflow RAG 대응 설계를 제안하였다. 어떤 자료가 어떤 변경을 지지했고 그 관계가 확정 비중까지 남았는지, 공통 오류를 고친 뒤 관련 판단도 다시 검토하는지를 함께 확인하도록 구성하였다. 같은 검색 근거의 단계별 배치와 전체 적격 원문 조건의 비교, 실제 펀드 운용조건과 수리 기준배분을 사용하는 확장 검증을 거쳐야 완화효과를 평가할 수 있다. 본실험과 대응 설계를 통해 제시한 핵심은 근거 제공 시점, 검토 절차와 단계 간 정보 전달을 함께 설계하여 필요한 차이와 판단 품질을 보존해야 한다는 것이다.')])
    rename = {
        '4 예비실험의 자료와 설계':'4 자료와 본실험 설계', '5 예비실험 평가방법':'5 평가방법', '6 예비실험 결과':'6 본실험 결과',
        '6 5 예비 질문에 대한 답과 본실험의 과제':'6 5 연구질문에 대한 답',
        '7 예비결과와 Workflow RAG 설계의 연결':'7 본실험 결과와 Workflow RAG 설계의 연결',
        '8 2 계획 단계: 근거와 운용조건을 구분하여 유지':'8 2 계획 단계의 근거와 운용조건 유지',
        '8 3 배분안 구성 단계: 필요한 차이를 실행 가능한 변경으로 연결':'8 3 배분안 구성 단계의 근거와 변경분 연결',
        '8 4 점검과 수정 단계: 설명을 확정 비중과 대조':'8 4 점검과 수정 단계의 설명과 확정 비중 대조',
        '9 한계와 본실험의 검증 계획':'9 한계와 후속 검증',
        '9 1 예비관찰로 확인한 범위와 남은 질문':'9 1 본실험의 해석 범위와 한계',
        '9 2 근거 제공 단계와 정보 접근을 나누는 통제 비교':'9 2 Workflow RAG의 배치와 정보 접근 검증',
        '9 3 기준배분, 판단 품질과 새 문서에서의 재검증':'9 3 실제 펀드와 기준배분을 이용한 확장 검증'}
    for b in blocks:
        if b['type']=='heading': b['text'] = rename.get(b['text'], b['text'])
        elif b['type']=='paragraph':
            text = b['text']
            text = text.replace('본실험에서는 계획, 배분안 구성, 점검과 수정의 세 단계', '후속 검증에서는 계획, 배분안 구성, 점검과 수정의 세 단계')
            text = text.replace('본실험의 제출 형식을 조건 간 동일하게 고정', '후속 비교의 제출 형식을 조건 간 동일하게 고정')
            text = text.replace('본 실험의 UBS 입력에서', '본실험의 UBS 입력에서')
            text = text.replace('초기 출력을 본 뒤 완료된 예비실험을 주된 탐색 분석으로 사용하고 외부 수학 처리를 추가하였다.', '초기 출력을 본 뒤 완료된 실행을 본실험의 분석 대상으로 정리하고 외부 수학 처리를 추가하였다.')
            text = text.replace('이 사례가 보여주는 것은 공유한 초안의 오류와 서로 다른 최종 배분이 공존했다는 점이다.', '이 사례에서는 별도 요청에서 나온 동일한 완료 D0 내용의 오류가 서로 다른 최종 배분에도 남았다.')
            text = text.replace('본 연구는 이를 계획, 배분안 구성, 점검과 수정의 3단계 중 어디에 처음 배치할지 비교한다.', '본 연구는 이를 계획, 배분안 구성, 점검과 수정의 3단계 중 어디에 처음 배치할지 후속 검증에서 비교하도록 설계한다.')
            text = text.replace('본 연구의 Workflow RAG는 검색 자체의 유무와 더불어 같은 근거를 투자판단의 어느 단계에 처음 제공하고 이후 어떻게 유지할지를 비교한다.', '본 연구의 Workflow RAG 대응 설계는 검색의 유무와 함께 같은 근거의 최초 제공 단계 및 이후 유지 방식을 후속 검증에서 비교하도록 구성한다.')
            text = text.replace('같은 초안을 공유하는 조건', '완료 D0 내용이 동일한 조건')
            text = text.replace('본고는 문제 제기, 본실험, 결과 해석과 대응 설계를 하나의 과정으로 연결하고, 확인한 단절을 Workflow RAG의 근거 전달과 검토 절차에 반영한다.', '본고는 확인한 단절을 Workflow RAG의 근거 전달과 검토 절차에 반영한다.')
            text = text.replace('이는 공통 GenAI에 기관별 자료를 넣었을 때 판단의 차이가 어떤 근거를 갖는지 살펴보는 소논문의 실험 과정이다.', '이 비교를 통해 공통 GenAI에 기관별 자료를 넣었을 때 판단의 차이가 어떤 근거를 갖는지 살펴본다.')
            for term in ['예비분석','예비관찰','예비실험','예비과제','예비진단','예비단계','예비결과']:
                text = text.replace(term, '본실험')
            b['text'] = text
    assert not any('예비' in b.get('text','') for b in blocks)
    return m

def revise():
    old = json.loads(BASE.read_text(encoding='utf-8'))
    m = deepcopy(old)
    m.update(title=TITLE, subtitle=SUBTITLE, version='5.4', author='배라31 2팀')
    replacements = {
        1: [p(PDF_P1), p(PDF_P2),
            p('현재까지는 기관별 근거가 실제 선택으로 이어지는지 먼저 확인하기 위해, Fidelity와 UBS의 공개 산업전망을 직접 제공한 예비분석을 수행하였다. 동일한 시장 자료와 미국 주식 11개 산업의 새 자금 배분 과제에서 초기·후기 제공과 미제공의 다섯 조건을 비교하였다. 네 기관 자료 조건과 대조군의 파생 배분 거리는 4.76~12.98%였으나, 정확한 자료 요약 뒤에도 서로 다른 배분에 같은 수치 오류가 남고 설명과 실제 비중이 어긋났다. 외부 정규화가 원비중의 증감 방향을 뒤집는 사례도 확인하였다. 이는 배분 차이뿐 아니라 근거의 전달과 판단 품질을 함께 평가해야 함을 보여주는 예비결과다.'),
            p('후속 본실험은 미국 액티브 주식형 펀드 4개의 공개 보유자료, 투자설명서와 운용보고서를 바탕으로 펀드별 출발 조건을 구성하도록 설계한다. 산업 범위는 현재 예비과제에서 사용한 GICS 11개 체계에 맞추고, 목표 비중과 기존 보유 대비 변경분을 함께 다룬다. 허용된 정보와 명시 제약을 반영한 수리 기준배분을 계산하여 필요한 배분 차이의 비교 기준으로 삼는다. 대응방안인 검색증강생성(RAG)은 투자판단의 계획, 배분안 구성, 점검과 수정의 3단계에 배치한다. 같은 검색 근거의 최초 제공 단계만 달리하고 이후에는 유지하며, 추가 근거가 없는 조건과 전체 적격 원문을 처음부터 제공하는 조건도 비교한다.'),
            p('본실험에서는 기준배분 대비 펀드 간 차이, 개별 정보 변화에 대한 반응, 모형상 기대수익 손실과 투자제약 위반을 함께 평가할 계획이다. 실제문서의 정확성과 공동오류를 점검하고, 보류한 새 문서에서 검색을 다시 수행하여 근거 활용 효과의 유지 여부와 검색·실행 비용도 확인한다. 4개 펀드의 기준배분 비교와 3단계 RAG 배치·완화효과는 아직 검증하지 않았다. 본 연구는 예비결과를 반영한 이 설계를 통해 필요한 판단 차이와 품질이 함께 유지되는 조건을 찾고, 근거 제공 시점, 검토 절차와 단계 간 정보 전달에 관한 설계 원칙을 제안한다.')],
        2: [p('주제어  생성형 AI, 기관투자자, 전략적 이질성, 과잉수렴, 자산배분, Workflow RAG')],
        4: [
            p('본 연구의 질문은 공통 생성형 AI(GenAI)가 기관투자자의 전략적 이질성을 축소하는가이다. 액티브 자산운용에서 기관별 투자목표, 기존 보유자산과 위험제약에 맞는 판단을 유지하려면, 공통 AI가 정보 처리와 분석에 관여하는 과정부터 살펴야 한다. Morgan Stanley가 2024년 출시한 AskResearchGPT는 자체 리서치의 검색과 요약을 지원하는 사례다(Morgan Stanley, 2024). 영국은행·FCA의 같은 해 조사에서는 응답 금융회사의 75%가 AI를 사용했지만 전체 AI 활용사례 중 완전자율 의사결정은 2%였다(Bank of England & FCA, 2024). 이 조사는 AI 전반에 관한 것이며, 두 수치의 분모도 각각 응답 회사와 활용사례로 다르다.'),
            p('중요한 것은 최종 결정을 통째로 자동화했는지에만 있지 않다. 자료를 찾고 요약하며 전망을 분석하고 배분안의 이유를 작성하는 각각의 업무에서 AI가 선택한 정보와 해석은 다음 단계의 입력이 될 수 있다. 서로 다른 기관의 최종 승인자가 독립적으로 판단하더라도, 판단 이전에 검토할 대안과 위험의 범위가 같은 모델의 해석에 의존한다면 그 영향이 여러 단계에 걸쳐 누적될 가능성이 있다. 본고는 이러한 분산 사용을 공통 GenAI가 투자판단에 관여하는 발생경로로 본다.'),
            p('액티브 펀드의 판단이 서로 달라야 할 이유는 관찰 가능한 운용조건에 있다. 같은 산업을 긍정적으로 평가하더라도 해당 산업을 이미 많이 보유한 펀드와 적게 보유한 펀드가 선택할 추가 배분은 다를 수 있다. 투자목표와 허용 위험이 다른 경우에도 적절한 선택은 같지 않을 수 있다. 이러한 차이는 투자자의 요구에 맞는 운용과 전략의 차별화를 뒷받침한다. 공통 정보가 같은 결론을 지지하는 합리적 수렴과, 펀드별 조건상 필요한 차이가 줄어드는 과잉수렴을 구분해야 하는 이유다.'),
            p('본 연구는 개별 조건상 필요한 배분 차이가 기준보다 작아지는 현상을 과잉수렴으로 정의한다. 배분 간 거리가 작다는 사실만으로 이를 판정하지 않고, 해당 펀드에 허용된 정보와 명시 제약에 따른 기준배분 및 판단 품질과 함께 비교한다. 공통 알고리즘의 개별 성능과 집단 후생, 공유 모델과 결과 동질화의 관계를 나누는 선행연구도 이 구분을 뒷받침한다(Kleinberg와 Raghavan, 2021; Bommasani 등, 2022). 유사한 권고가 실제 거래와 위험 노출의 동조로 이어진다면 금융안정의 문제가 될 수 있으나, 그 시장 수준의 연결은 본 예비분석에서 관측하지 않았다(FSB, 2024).'),
            p('대응방안의 중심은 기관별 근거를 투자판단 과정에 연결하는 Workflow RAG다. RAG는 외부 자료를 검색하여 생성 과정에 연결하는 접근이며(Lewis 등, 2020), 본 연구는 이를 계획, 배분안 구성, 점검과 수정의 3단계 중 어디에 처음 배치할지 비교한다. 자료를 확보했는지와 그 자료가 최종 판단까지 유지되었는지는 서로 다른 질문이다. 따라서 같은 검색 결과를 어느 단계부터 제공하고 이후 어떻게 전달할지, 정보가 들어온 뒤 어떤 판단을 다시 검토할지까지 함께 설계한다.'),
            p('전체 본실험은 미국 액티브 주식형 펀드 4개의 목표·보유·제약을 구성하고, 같은 모델이 제시한 배분을 조건부 수리 기준배분과 비교하도록 설계한다. 이 연구에 앞서 현재까지 완료한 것은 두 기관의 공개 산업전망을 직접 제공한 다섯 조건의 예비분석이다. 새 자금을 GICS 11개 산업에 배분하게 하여 기관 자료의 의미, 모델이 밝힌 이유와 실제 비중을 대조하였다. 실제 펀드의 보유조건을 반영한 기준배분이나 검색 RAG의 완화효과를 측정한 단계와는 구분된다.'),
            p('예비분석에서는 기관 자료 조건의 배분이 대조군과 달랐어도 설명과 실제 선택이 어긋났고, 서로 다른 배분에 같은 수치 오류가 남았다. 원비중과 외부 정규화된 배분의 증감 방향도 일치하지 않는 경우가 있었다. 본고는 이 결과를 원래 질문에 대한 확정 답으로 쓰기보다, 전략적 이질성과 품질을 함께 평가하려면 근거의 단계 간 전달과 확정 비중까지 확인해야 한다는 설계 근거로 사용한다. 관측된 실패를 3단계 Workflow RAG의 검사 항목에 연결하고, 필요한 차이와 품질이 함께 유지되는 조건을 검증할 본실험을 제시하는 것이 현 단계의 기여다.')
        ],
        5: [], 6: [], 7: [], 8: [], 9: [],
        21: [
            p('전체 연구질문 RQ1은 동일한 GenAI가 펀드별 운용조건상 필요한 배분 차이를 축소하고, 그 축소가 판단 손실이나 투자제약 위반을 동반하는가이다. 필요한 차이는 같은 통제 과제의 허용 정보와 명시 제약을 반영한 기준배분으로 정의한다. 단순한 출력 다양성과 기관에 적합한 판단의 차이를 구분하는 질문이다.'),
            p('RQ2는 펀드별 근거를 Workflow RAG의 어느 단계부터 제공·유지할 때 필요한 판단 차이와 품질이 함께 보존되는가이다. 계획, 배분안 구성, 점검과 수정의 최초 제공 조건을 비교하고, 무추가 근거 및 전체 적격 원문 조건으로 정보 접근 범위와 단계 배치의 역할을 나누어 살펴본다.'),
            p('현재 예비분석은 본실험에 앞서 자료가 선택에 연결되는지를 살펴보는 세 하위 질문 P1~P3을 다룬다. 아래 4~6장의 결과는 이 하위 질문에 대응하며, 전체 RQ1·RQ2의 완화효과 검증은 8~9장의 후속 설계로 제시한다.')],
        22: [p('P1은 기관별 공개 산업전망의 의미가 설명과 실제 비중 사이에서 유지되는가이다. 원문의 방향·범위·이유를 산업행 요약과 선택 이유에 대조하고, 모델이 밝힌 이유가 제출 비중과 일치하는지 확인한다. 기관 전망을 따르지 않았더라도 원문을 정확히 이해하고 별도 판단의 이유를 밝혔으면 그 자체를 오독으로 분류하지 않는다.')],
        23: [p('P2는 기관 전망을 제공한 조건의 최종 배분이 공통 시장 자료만 제공한 대조군과 어떻게 다른가이다. 같은 외부 계산을 적용한 배분 사이의 거리를 비교한다. 이 값은 결과의 차이를 보여주며, 필요한 전략적 차이나 과잉수렴의 크기를 뜻하지 않는다.')],
        24: [p('P3은 기관 자료를 초안부터 제공한 경우와 최종 검토 때 처음 제공한 경우에 근거 해석과 배분 수정이 어떻게 다른가이다. 같은 기관의 초기·후기 조건 및 각 조건의 초안·최종안을 비교한다. 자료 제공과 함께 초안·노출 횟수·대화 이력도 달라지므로, 이 비교는 순수한 시점 효과나 RAG 배치 효과의 검증이 아니다.')],
        63: [p('아래 결과는 전체 연구의 RQ1·RQ2에 앞선 예비 질문 P1~P3에 대응한다. 기관 전망을 제공한 네 조건의 최종 배분은 대조군과 달랐다. 그 차이가 기관 근거의 충실한 반영에서 나왔는지 확인하기 위해 원비중의 유효성, 최종 배분의 차이와 설명·숫자의 대응을 함께 살펴본다.')],
        80: [p('P1에서는 기관별 산업행의 등급과 요약이 보존되었으나 이후 해석과 실제 비중의 연결은 일관되지 않았다. P2에서는 네 기관 자료 조건의 파생 배분이 C와 모두 달랐다. P3에서는 초기·후기 조건의 수정 양상이 달랐지만 어느 절차가 근거를 더 충실히 반영했는지는 확정하지 못하였다. 따라서 현재 관측으로 과잉수렴의 존재나 Workflow RAG의 완화효과를 판정할 수는 없다. 예비결과는 본실험에서 배분 거리와 함께 근거 반응, 공동오류와 판단 품질을 검토해야 한다는 근거를 제공한다.')],
        82: [p('예비결과의 의미는 원래 연구질문을 대신하는 데 있지 않다. 기관별 자료가 요약에는 반영되어도 최종 선택의 근거로 유지되지 않을 수 있다는 관측은, Workflow RAG를 배치할 때 검색 결과를 제공한 사실과 판단에 사용한 결과를 나누어 확인해야 함을 보여준다. 다음 해석은 이 연결이 끊긴 지점을 계획·배분안 구성·점검과 수정의 설계에 연결한다.')],
        138: [p('본 연구는 공통 GenAI가 기관투자자의 전략적 이질성을 축소하는가라는 질문에서 출발하여, 필요한 판단 차이와 품질을 함께 유지할 Workflow RAG의 배치를 연구한다. 액티브 펀드의 목표·기존 보유·위험제약이 다른데도 공통 모델의 해석에 의존하여 필요한 배분 차이가 줄어드는지, 펀드별 근거를 어느 단계부터 제공·유지해야 이를 완화할 수 있는지가 전체 연구의 중심이다.'),
            p('현재 완료한 공개 산업전망 예비분석에서는 다섯 조건의 배분 차이와 함께 정확한 요약 뒤의 설명·비중 불일치, 서로 다른 배분에 남은 공유 수치 오류, 외부 정규화에 따른 증감 방향 역전을 확인하였다. 이 결과만으로 전략적 이질성의 축소나 RAG의 개선 효과를 판정할 수는 없다. 다만 근거를 검색·제공하는 것에 더해 그 근거가 어떤 판단과 변경을 지지했고 실제 확정 비중에 어떻게 남았는지 확인해야 한다는 설계상 요구를 구체화하였다.')],
        139: [p('이에 본고는 계획, 배분안 구성, 점검과 수정의 3단계에 같은 검색 근거를 처음 투입하고 이후 유지하는 본실험을 제시하였다. 미국 액티브 주식형 펀드 4개의 운용조건과 수리 기준배분을 구성한 뒤, 근거 미제공·단계별 제공·전체 적격 원문 제공 조건에서 필요한 차이, 모형상 판단 손실과 제약 위반을 함께 평가해야 한다. 보류한 새 문서에서 검색을 다시 수행하고 비용도 확인해야 완화효과를 평가할 수 있다. 현재의 기여는 이 전체 컨셉을 대체하는 오류 목록이 아니라, 예비관찰을 바탕으로 근거 제공 시점과 검토 절차, 단계 간 정보 전달을 함께 검증할 설계를 구체화한 데 있다.')],
        185: [p(old['blocks'][185]['text'].replace('본 원고 v5.3', '본 원고 v5.4'))]
    }
    heading_changes = {
        '4 자료와 실험설계': '4 예비실험의 자료와 설계',
        '5 평가방법': '5 예비실험 평가방법',
        '6 실험 결과': '6 예비실험 결과',
        '6 5 연구질문에 대한 답': '6 5 예비 질문에 대한 답과 본실험의 과제',
        '7 실험 결과의 재해석': '7 예비결과와 Workflow RAG 설계의 연결',
    }
    insert_after = {
        16: [p('RAG는 외부 문서를 검색하여 모델의 생성 과정에 연결하는 접근이다(Lewis 등, 2020). 본 연구의 Workflow RAG는 검색 자체의 유무와 더불어 같은 근거를 투자판단의 어느 단계에 처음 제공하고 이후 어떻게 유지할지를 비교한다. 검색 정확성, 제공한 자료의 범위, 단계별 근거 활용은 구분하여 평가해야 한다. 현재의 직접 제공 예비실험은 마지막 연결에서 나타날 수 있는 실패를 확인하는 역할을 한다.')],
        25: [p('전체 연구의 본실험은 실제 펀드의 운용조건과 수리 기준배분을 포함하도록 설계한다. 아래에서는 그에 앞서 완료한 두 공개 전망·다섯 조건의 예비실험을 기술한다. 예비과제는 현재 정리된 GICS 최상위 11개 산업과 새 자금 배분을 사용하였으며, 실제 보유자산의 조정을 실행한 과제는 아니다. 후속 4개 펀드의 자료 구성과 기준 비교는 8.6 및 9장에서 구체화한다.')],
        46: [p('예비실험에서는 제공 문서와 설명·비중의 대응을 평가한다. 전체 본실험에서 사용할 펀드별 기준배분 대비 손실과 과잉수렴은 아직 계산하지 않았다. 두 평가를 구분하여, 현재의 산출물 차이를 필요한 전략적 차이의 회복으로 오인하지 않도록 한다.')]
    }
    workflow = json.loads((WORK / 'workflow_rag_draft.json').read_text(encoding='utf-8'))['blocks']
    assert all('??' not in b.get('text', '') for b in workflow), 'Workflow draft contains corrupt text'
    blocks = []
    for i, block in enumerate(old['blocks']):
        if 99 <= i < 137:
            if i == 99: blocks.extend(deepcopy(workflow))
            continue
        b = deepcopy(block)
        if b['type'] == 'heading': b['text'] = heading_changes.get(b['text'], b['text'])
        blocks.extend(deepcopy(replacements.get(i, [b])))
        blocks.extend(deepcopy(insert_after.get(i, [])))
    references = [
        {'type': 'reference', 'id': 'boe_fca2024', 'text': 'Bank of England, & Financial Conduct Authority. (2024, November 21). Artificial intelligence in UK financial services – 2024.', 'url': 'https://www.bankofengland.co.uk/report/2024/artificial-intelligence-in-uk-financial-services-2024'},
        {'type': 'reference', 'id': 'morgan2024', 'text': 'Morgan Stanley. (2024, October 23). Morgan Stanley Research announces AskResearchGPT powered by OpenAI.', 'url': 'https://www.morganstanley.com/press-releases/morgan-stanley-research-announces-askresearchgpt'},
        {'type': 'reference', 'id': 'lewis2020', 'text': 'Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Küttler, H., Lewis, M., Yih, W.-t., Rocktäschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. Advances in Neural Information Processing Systems, 33, 9459–9474.', 'url': 'https://proceedings.neurips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html'}
    ]
    at = next(i for i,b in enumerate(blocks) if b['type']=='heading' and b['text']=='자료 출처')
    blocks[at:at] = references
    m['blocks'] = blocks
    m = reframe_as_main_experiment(m, old)
    blocks = m['blocks']
    assert [b for b in blocks if b['type']=='table'] == [b for b in old['blocks'] if b['type']=='table']
    old_refs = [b for b in old['blocks'] if b['type']=='reference']
    assert all(b in blocks for b in old_refs)
    assert '??' not in json.dumps(m, ensure_ascii=False)
    OUT.write_text(json.dumps(m, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    publisher = module('libera_publish_540', ROOT / 'research/manuscript_v520_20261026/publish_manuscript.py')
    publisher.render(OUT, ROOT / 'paper' / (STEM+'.md'), ROOT / 'paper' / (STEM+'.html'))
    return m

def build(m):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    builder = module('libera_word_540', ROOT / 'research/document_v520_20261026/build_document.py')
    builder.SOURCE, builder.WORK = OUT, WORK
    builder.OUT = ROOT / 'paper' / (STEM+'.docx')
    original_add_table = builder.add_table
    def add_table(doc, b):
        original_add_table(doc, b)
        if b['caption'].startswith('표 5 '):
            for row in doc.tables[-1].rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs: paragraph.paragraph_format.keep_with_next = True
    builder.add_table = add_table
    figures = json.loads(builder.MANIFEST.read_text(encoding='utf-8'))
    if isinstance(figures, dict) and 'figures' in figures: figures = figures['figures']
    figures = {b['id']:b for b in figures} if isinstance(figures, list) else figures
    ill = builder.compose(m, figures)
    ill['version'] = '5.4'
    for index,b in enumerate(ill['blocks']):
        if b.get('id')=='fig01_allocations':
            ill['blocks'][index-1:index+1] = [b, ill['blocks'][index-1]]
            break
    for index,b in enumerate(ill['blocks']):
        if b.get('id')=='fig04_normalization':
            fig = ill['blocks'].pop(index)
            following = next(i for i in range(index,len(ill['blocks'])) if ill['blocks'][i]['type']=='heading')
            ill['blocks'].insert(following, fig)
            break
    design = json.loads((WORK / 'figures/manifest.json').read_text(encoding='utf-8'))[0]
    design.update(type='figure', display_caption='그림 5 본실험의 관측을 반영한 Workflow RAG 대응 설계')
    idx = next(i for i,b in enumerate(ill['blocks']) if b['type']=='heading' and b['text'].startswith('8 1 '))
    ill['blocks'].insert(idx, design)
    ILL = WORK / 'illustrated_manuscript.json'
    ILL.write_text(json.dumps(ill, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    doc = builder.make_doc(ill)
    for para in doc.paragraphs[:2]: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author = doc.paragraphs[2].insert_paragraph_before(m['author'])
    author.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    author.paragraph_format.keep_with_next = True
    author.paragraph_format.space_after = Pt(4)
    for run in author.runs: run.font.size = Pt(9.5)
    doc.core_properties.author = m['author']
    doc.save(builder.OUT)
    result = {'source':str(OUT), 'source_sha256':digest(OUT), 'base_source_sha256':digest(BASE), 'illustrated_source_sha256':digest(ILL), 'docx':str(builder.OUT), 'docx_sha256':digest(builder.OUT), 'title':TITLE, 'subtitle':SUBTITLE, 'tables':len(doc.tables), 'figures':len(doc.inline_shapes), 'native_equations':2, 'source_pdf': 'C:/Users/Insun/Downloads/greengreen (1).pdf', 'source_pdf_sha256':digest(Path('C:/Users/Insun/Downloads/greengreen (1).pdf'))}
    (WORK / 'build_receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__=='__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('--source-only', action='store_true'); args = ap.parse_args()
    manuscript = revise()
    if not args.source_only: build(manuscript)
