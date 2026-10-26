"""Build the illustrated v5.2 paper with native Word tables and equations."""
from pathlib import Path
from decimal import Decimal
from copy import deepcopy
import hashlib
import json
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
SOURCE = ROOT / 'paper/v520_manuscript_20261026.json'
DATA = ROOT / 'research/last_final_20261026/derived_allocation/completed_five_condition_dataset.json'
OUT = ROOT / 'paper/LIBERA_소논문_시각화포함_v5.2_20261026.docx'
MANIFEST = WORK / 'figures/manifest.json'


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def node(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def math_run(text, upright=False):
    r = node('m:r')
    if upright:
        prop = node('m:rPr')
        prop.append(node('m:sty', **{'m:val': 'p'}))
        r.append(prop)
    t = node('m:t')
    t.text = text
    r.append(t)
    return r


def scripted(kind, base, value):
    e = node('m:' + kind)
    b = node('m:e'); b.append(math_run(base)); e.append(b)
    s = node('m:sub' if kind == 'sSub' else 'm:sup')
    s.append(math_run(value, value == 'ann'))
    e.append(s)
    return e


def add_equation(doc, kind, number):
    p = doc.add_paragraph(style='Equation')
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    math = node('m:oMath')
    if kind == 'risk':
        math.append(scripted('sSub', 'σ', 'ann'))
        math.append(math_run('(w) = '))
        rad = node('m:rad')
        pr = node('m:radPr'); pr.append(node('m:degHide', **{'m:val': 1})); rad.append(pr)
        rad.append(node('m:deg'))
        content = node('m:e')
        content.append(math_run('252 '))
        content.append(scripted('sSup', 'w', 'T'))
        content.append(math_run('Σw'))
        rad.append(content); math.append(rad)
    else:
        math.append(math_run('TV(p,q) = ', True))
        f = node('m:f')
        num = node('m:num'); num.append(math_run('1')); f.append(num)
        den = node('m:den'); den.append(math_run('2')); f.append(den)
        math.append(f)
        summation = node('m:nary')
        prop = node('m:naryPr')
        prop.append(node('m:chr', **{'m:val': '∑'}))
        prop.append(node('m:limLoc', **{'m:val': 'undOvr'}))
        summation.append(prop)
        sub = node('m:sub'); sub.append(math_run('i=1')); summation.append(sub)
        sup = node('m:sup'); sup.append(math_run('11')); summation.append(sup)
        content = node('m:e')
        delim = node('m:d')
        prop = node('m:dPr')
        prop.append(node('m:begChr', **{'m:val': '|'}))
        prop.append(node('m:endChr', **{'m:val': '|'})); delim.append(prop)
        inner = node('m:e')
        inner.append(scripted('sSub', 'p', 'i'))
        inner.append(math_run('−'))
        inner.append(scripted('sSub', 'q', 'i'))
        delim.append(inner); content.append(delim)
        summation.append(content); math.append(summation)
    p._p.append(math)
    p.add_run('\t(' + str(number) + ')')


def set_font(style, size, bold=False):
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + name), 'Malgun Gothic')
    for borders in style.element.xpath('.//w:pBdr'):
        borders.getparent().remove(borders)
    for c in style.element.xpath('.//w:color'):
        c.attrib.clear(); c.set(qn('w:val'), '000000')


def add_link(p, label, url):
    if not re.match(r'^[A-Za-z][A-Za-z0-9+.-]*:', url):
        url = (OUT.parent / url).resolve().as_uri()
    h = node('w:hyperlink', **{'r:id': p.part.relate_to(url, RT.HYPERLINK, is_external=True)})
    r = node('w:r')
    props = node('w:rPr')
    props.append(node('w:color', **{'w:val': '000000'}))
    props.append(node('w:u', **{'w:val': 'single'})); r.append(props)
    t = node('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def add_table(doc, b):
    caption = doc.add_paragraph(b['caption'], 'Caption')
    caption.paragraph_format.keep_with_next = True
    widths = b.get('widths', [16 / len(b['headers'])] * len(b['headers']))
    widths = [w * 16 / sum(widths) for w in widths]
    t = doc.add_table(rows=1, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    borders = node('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(node('w:' + edge, **{'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
    t._tbl.tblPr.append(borders)
    margins = node('w:tblCellMar')
    for side, value in [('top', 75), ('bottom', 75), ('left', 95), ('right', 95)]:
        margins.append(node('w:' + side, **{'w:w': value, 'w:type': 'dxa'}))
    t._tbl.tblPr.append(margins)
    for c, value in zip(t.rows[0].cells, b['headers']):
        c.text = str(value)
    t.rows[0]._tr.get_or_add_trPr().append(node('w:tblHeader'))
    for row in b['rows']:
        assert len(row) == len(widths)
        for c, value in zip(t.add_row().cells, row):
            c.text = str(value)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(node('w:cantSplit'))
        for ci, c in enumerate(row.cells):
            c.width = Cm(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                c._tc.get_or_add_tcPr().append(node('w:shd', **{'w:fill': 'E8EDF1'}))
            for p in c.paragraphs:
                p.style = doc.styles['Table Text']
                p.paragraph_format.keep_with_next = ri == 0
                p.paragraph_format.keep_together = True
                if ri == 0 or re.fullmatch(r'[\d.,%+−\- ]+', c.text):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = ri == 0
    if b.get('note'):
        doc.add_paragraph(b['note'], 'Table Note')


def figure_path(b):
    p = Path(b['png'])
    return p if p.is_absolute() else WORK / 'figures' / p


def add_figure(doc, b):
    p = doc.add_paragraph(style='Figure')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    picture = p.add_run().add_picture(str(figure_path(b)), width=Cm(16))
    picture._inline.docPr.set('descr', b['title'] + ' ' + b.get('caption', ''))
    cap = doc.add_paragraph(b['display_caption'], 'Caption')
    cap.paragraph_format.keep_with_next = bool(b.get('note'))
    if b.get('note'):
        doc.add_paragraph(b['note'], 'Figure Note')


def compose(source, figures):
    m = deepcopy(source)
    tables = {int(b['caption'].split()[1]): deepcopy(b) for b in m['blocks'] if b['type'] == 'table'}
    fig_ids = ['fig01_allocations', 'fig02_distances', 'fig03_evidence', 'fig04_normalization', 'figA01_control_difference']
    figs = {}
    for i, key in enumerate(fig_ids):
        b = deepcopy(figures[key])
        b['type'] = 'figure'
        number = str(i + 1) if i < 4 else 'A1'
        titles = ['조건별 최종 파생 산업배분', '비교 단위별 배분 거리와 초안 이후 수정거리', '자료 재현 이후 설명과 선택의 불일치 사례', '정규화에 따른 정보기술 비중의 증감 방향 역전', '공통 대조군 대비 최종 파생 비중의 차이']
        b['display_caption'] = '그림 ' + number + ' ' + titles[i]
        if key == 'fig02_distances':
            b['note'] = b['note'].replace('거리는 총변동거리 ½Σ|wᵢ−vᵢ|를 백분율로 표시했다.', '거리는 식 (2)의 TV 거리를 전체 자금의 백분율로 표시한 값이다.')
        figs[key] = b

    def paragraph(text):
        return {'type': 'paragraph', 'text': text}

    out = []
    current = ''
    figure4_added = False
    for b in m['blocks']:
        if b['type'] == 'heading':
            # Figure 4 follows the first analytical paragraph of 7.3 below.
            current = b['text']
        if b['type'] == 'table':
            n = int(b['caption'].split()[1])
            if n in (6, 7, 8):
                out.append(figs[fig_ids[n - 6]])
                continue
            if n in (9, 10, 11):
                b['caption'] = b['caption'].replace('표 ' + str(n) + ' ', '표 ' + str(n - 3) + ' ', 1)
        if b['type'] == 'paragraph':
            txt = b['text']
            if current == '6 2 다섯 조건의 최종 산업배분':
                txt = txt.replace('(표 6)', '(그림 1; 부록 표 A1)')
            elif current == '6 3 기관 자료 및 제공 절차별 배분 차이':
                txt = txt.replace('(표 7)', '(그림 2; 부록 표 A2)')
            elif current == '6 4 자료 요약 이후에 나타난 근거와 수치의 단절':
                txt = txt.replace('표 8은 세 유형의 대표 사례를 원문 및 계산값과 대조한 결과이다(E2).', '그림 3은 등급·수치·선택 서술이 어긋난 대표 사례를 대조한다. 대표 사례의 상세 대조와 추정 표현의 판정은 부록 표 A3에 제시하였다(E2).')
                txt = txt.replace('표의 공유 수치 오류는', '그림 3과 부록 표 A3의 공유 수치 오류는')
            elif current == '7 1 정확한 자료 재현과 실제 판단의 분리':
                txt = txt.replace('표 8의 사례는', '그림 3과 부록 표 A3의 사례는')
            for old, new in [(9, 6), (10, 7), (11, 8)]:
                txt = re.sub(r'표 ' + str(old) + r'(?=\D|$)', '표 ' + str(new), txt)
            txt = txt.replace('표 6와 같이', '표 6과 같이')
            if '√(252wᵀΣw)' in txt:
                txt = txt.replace('비중 벡터 w와 일별 공분산 Σ에 대해 연율 변동성은 √(252wᵀΣw)로 계산한다.', '비중 벡터 w와 일별 공분산 Σ에 대한 연율 변동성은 식 (1)과 같이 계산한다.')
                b['text'] = txt; out.append(b)
                out.append({'type': 'equation', 'kind': 'risk', 'number': 1})
                continue
            if 'TV(p,q)=½Σᵢ|pᵢ−qᵢ|' in txt:
                txt = txt.replace('합계가 1인 두 파생 비중 벡터 p와 q의 TV 거리는 TV(p,q)=½Σᵢ|pᵢ−qᵢ|로 정의한다.', '합계가 1인 두 파생 비중 벡터 p와 q의 TV 거리는 식 (2)로 정의한다.')
                b['text'] = txt; out.append(b)
                out.append({'type': 'equation', 'kind': 'tv', 'number': 2})
                continue
            if current == '7 3 배분의 방향을 바꿀 수 있는 정규화' and not figure4_added:
                txt += ' 그림 4는 이 사후 진단을 원비중과 파생 비중으로 나누어 보여준다.'
                b['text'] = txt; out.append(b)
                out.append(figs['fig04_normalization']); figure4_added = True
                continue
            b['text'] = txt
        out.append(b)

    # The main manuscript stays intact; numeric supplements follow its evidence appendix.
    out.append({'type': 'heading', 'text': '부록 수치표와 추가 시각화', 'level': 1})
    out.append(paragraph('본문 그림의 정확한 값과 사례 판정의 범위를 확인할 수 있도록 전체 수치표를 제시한다. 표 A1~A3은 시각화에 대응하는 기존 결과표이며, 표 A4~A8은 각 조건의 초안 D0와 첫 최종안 F0에서 원비중과 파생 비중을 함께 보고한다. 이 행들은 독립 실험의 수를 뜻하지 않는다.'))
    for n in (6, 7, 8):
        b = tables[n]
        b['caption'] = re.sub(r'^표 \d+ ', '표 A' + str(n - 5) + ' ', b['caption'])
        out.append(b)
    dataset = json.loads(DATA.read_text(encoding='utf-8'))
    records = {(r['case_id'], r['stage']): r for r in dataset['stage_records']}
    names = ['Fidelity E', 'Fidelity L', 'UBS E', 'UBS L', 'C']
    sectors = [(row[0], row[1]) for row in tables[1]['rows']]
    for ordinal, (case, label) in enumerate(zip(dataset['case_order'], names), 4):
        d0, f0 = records[(case, 'D0')], records[(case, 'F0')]
        lookup0 = {r['sector_code']: r for r in d0['rows']}
        lookup1 = {r['sector_code']: r for r in f0['rows']}
        rows = []
        for code, name in sectors:
            a, b = lookup0[code], lookup1[code]
            rows.append([name, f"{Decimal(a['raw_weight_exact_decimal']) * 100:.2f}", a['final_percent'], f"{Decimal(b['raw_weight_exact_decimal']) * 100:.2f}", b['final_percent']])
        rows.append(['합계', f"{Decimal(d0['raw_sum_exact_decimal']) * 100:.2f}", '100.00', f"{Decimal(f0['raw_sum_exact_decimal']) * 100:.2f}", '100.00'])
        out.append({'type': 'table', 'caption': '표 A' + str(ordinal) + ' ' + label + ' 원비중과 파생 비중', 'headers': ['산업', 'D0 원비중', 'D0 파생 비중', 'F0 원비중', 'F0 파생 비중'], 'rows': rows, 'widths': [4.4, 2.9, 2.9, 2.9, 2.9], 'note': '단위 %. 파생값은 동일한 외부 계산 결과이며 원응답의 합계 오류를 소급해 제거하지 않는다. 출처 E1.'})
    out.append(figs['figA01_control_difference'])
    out.append(paragraph('그림 A1은 F0 파생 비중에서 같은 산업의 C 비중을 뺀 값이다. 양수는 C보다 높은 비중을, 음수는 낮은 비중을 뜻한다. C는 같은 시장 자료를 해석한 비교 조건이며 기관의 정답 배분은 아니다.'))
    m['blocks'] = out
    m['version'] = '5.2 도표 포함본'
    return m


def make_doc(m):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    sec.footer_distance = Cm(0.9)
    sizes = [('Normal', 10.5, False), ('Title', 19, True), ('Subtitle', 11.5, False), ('Heading 1', 13.5, True), ('Heading 2', 11, True), ('Caption', 9.5, True)]
    for name, size, bold in sizes:
        set_font(doc.styles[name], size, bold)
    for name, size in [('Table Text', 9), ('Table Note', 8.5), ('Figure Note', 8.5), ('Reference', 9), ('Figure', 10.5), ('Equation', 11)]:
        style = doc.styles.add_style(name, 1)
        style.base_style = doc.styles['Normal']; set_font(style, size)
    pf = doc.styles['Normal'].paragraph_format
    pf.line_spacing = 1.32; pf.space_after = Pt(6); pf.widow_control = True
    for name in ('Heading 1', 'Heading 2'):
        pf = doc.styles[name].paragraph_format
        pf.keep_with_next = True; pf.keep_together = True
        pf.space_before = Pt(13 if name == 'Heading 1' else 10); pf.space_after = Pt(6)
    for name in ('Table Text', 'Table Note', 'Figure Note', 'Reference'):
        pf = doc.styles[name].paragraph_format
        pf.line_spacing = 1.12
        pf.space_after = Pt(0 if name == 'Table Text' else 5)
        pf.first_line_indent = Cm(0)
    for name in ('Table Note', 'Figure Note'):
        doc.styles[name].paragraph_format.space_before = Pt(4)
    doc.styles['Reference'].paragraph_format.left_indent = Cm(0.5)
    doc.styles['Reference'].paragraph_format.first_line_indent = Cm(-0.5)
    doc.styles['Caption'].paragraph_format.space_before = Pt(7)
    doc.styles['Caption'].paragraph_format.space_after = Pt(5)
    doc.styles['Figure'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Figure'].paragraph_format.space_before = Pt(5)
    doc.styles['Figure'].paragraph_format.space_after = Pt(0)
    doc.styles['Equation'].paragraph_format.space_before = Pt(5)
    doc.styles['Equation'].paragraph_format.space_after = Pt(8)
    doc.add_paragraph(m['title'], 'Title')
    doc.add_paragraph(m['subtitle'], 'Subtitle')
    meta = doc.add_paragraph(m['date'] + '  |  v' + m['version'])
    meta.paragraph_format.space_after = Pt(9)
    for r in meta.runs:
        r.font.size = Pt(9)
    for b in m['blocks']:
        typ = b['type']
        if typ == 'heading':
            p = doc.add_paragraph(b['text'], 'Heading ' + str(b['level']))
        elif typ == 'paragraph':
            doc.add_paragraph(b['text'])
        elif typ == 'table':
            add_table(doc, b)
        elif typ == 'figure':
            add_figure(doc, b)
        elif typ == 'equation':
            add_equation(doc, b['kind'], b['number'])
        elif typ == 'reference':
            p = doc.add_paragraph(b['text'] + ' ', 'Reference')
            if b.get('url'):
                add_link(p, '원문', b['url'])
            for link in b.get('links', []):
                p.add_run('  '); add_link(p, link['label'], link['url'])
        else:
            raise ValueError(typ)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(); r.font.size = Pt(9)
    field = node('w:fldSimple', **{'w:instr': 'PAGE'})
    r._r.addnext(field)
    settings = doc.settings.element
    settings.append(node('w:updateFields', **{'w:val': 'true'}))
    doc.core_properties.title = m['title']
    doc.core_properties.subject = m['subtitle']
    doc.core_properties.author = ''
    doc.core_properties.comments = ''
    doc.core_properties.keywords = 'LIBERA, 생성형 AI, 산업배분, 근거 충실성'
    doc.save(OUT)
    return doc


def main():
    data = json.loads(MANIFEST.read_text(encoding='utf-8'))
    if isinstance(data, dict) and 'figures' in data:
        data = data['figures']
    figures = {b['id']: b for b in data} if isinstance(data, list) else data
    source = json.loads(SOURCE.read_text(encoding='utf-8'))
    m = compose(source, figures)
    src = WORK / 'illustrated_manuscript.json'
    src.write_text(json.dumps(m, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    doc = make_doc(m)
    result = {'input': str(SOURCE), 'input_sha256': digest(SOURCE), 'data_sha256': digest(DATA), 'illustrated_source': str(src), 'illustrated_source_sha256': digest(src), 'docx': str(OUT), 'docx_sha256': digest(OUT), 'tables': len(doc.tables), 'images': len(doc.inline_shapes), 'native_equations': 2, 'paragraphs': len(doc.paragraphs), 'runtime': sys.executable}
    (WORK / 'build_receipt.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
